from pathlib import Path
import sys

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


def iter_blocks(parent):
    parent_elm = parent.element.body if isinstance(parent, DocumentType) else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def main() -> None:
    source = Path(sys.argv[1])
    doc = Document(source)
    print(f"SOURCE: {source}")
    print(f"PARAGRAPHS: {len(doc.paragraphs)}")
    print(f"TABLES: {len(doc.tables)}")
    print("\nCONTENT\n")

    table_index = 0
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                style = block.style.name if block.style is not None else ""
                print(f"[P|{style}] {text}")
        else:
            table_index += 1
            print(f"[TABLE {table_index}]")
            for row in block.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                print(" | ".join(cells))
            print(f"[/TABLE {table_index}]")


if __name__ == "__main__":
    main()
