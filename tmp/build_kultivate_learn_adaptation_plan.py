from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\Joshua\OneDrive\Documents\Kultivate\outputs\Kultivate_Learn_Rencana_Adaptasi_Draft_v0.1_2026-09-10.docx")

NAVY = "102A3A"
BLUE = "1477A8"
PALE_BLUE = "EAF4F8"
PALE_GRAY = "F3F5F7"
MID_GRAY = "64727D"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=11, color=BLACK, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def add_text(paragraph, text, bold=False, italic=False, color=BLACK, size=11):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return run


def add_para(doc, text="", style=None, before=0, after=7, line=1.16, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    return p


def add_lead_para(doc, label, text, before=0, after=8):
    p = add_para(doc, before=before, after=after)
    add_text(p, label, bold=True)
    add_text(p, text)
    return p


def add_bullets(doc, items, level=0, font_size=10.7):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        add_text(p, item, size=font_size)
        p.paragraph_format.left_indent = Inches(0.22 + level * 0.22)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(4.5)
        p.paragraph_format.line_spacing = 1.12


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        add_text(p, item, size=10.7)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.12


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_text(p, text, bold=True, size=18 if level == 1 else 13.5)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    return p


def add_table(doc, headers, rows, widths, header_fill=NAVY, alternate=True, formatter_fn=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 120, 120, 120, 120)
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        add_text(p, value, bold=True, color=WHITE, size=9.3)

    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        if alternate and row_index % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, PALE_GRAY)
        for col_index, value in enumerate(row_data):
            cell = cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 105, 115, 105, 115)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            align = WD_ALIGN_PARAGRAPH.LEFT
            if col_index == 0 and len(str(value)) <= 12:
                align = WD_ALIGN_PARAGRAPH.CENTER
            p.alignment = align
            add_text(p, str(value), size=9.15)
            if formatter_fn:
                formatter_fn(cell, row_index, col_index, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    add_text(p, "Kultivate Learn  |  Draft v0.1  |  ", size=8.5, color=MID_GRAY)
    run = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color=MID_GRAY)


def set_document_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.paragraph_format.space_after = Pt(8)

    for name, size in (("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.keep_with_next = True


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    set_document_styles(doc)

    doc.core_properties.title = "Rencana Adaptasi Kultivate Learn"
    doc.core_properties.subject = "Draft internal untuk adaptasi course dan funnel konsultasi"
    doc.core_properties.author = "Kultivate"
    doc.core_properties.comments = "Draft v0.1 untuk revisi internal"

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    add_text(p, "KULTIVATE  /  LEARN", bold=True, size=10, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(title, "Rencana Adaptasi Kultivate Learn", bold=True, size=30)

    subtitle = add_para(doc, "Draft kurikulum, funnel, data, dan pelaksanaan untuk revisi internal", after=6)
    subtitle.runs[0].font.size = Pt(13)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    meta = add_para(doc, "Versi 0.1  |  10 September 2026", after=20)
    meta.runs[0].font.size = Pt(9.5)
    meta.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)

    add_heading(doc, "Keputusan utama", level=1)
    add_para(
        doc,
        "Kultivate Learn sebaiknya tidak langsung diperluas menjadi enam track dan 36 lesson. Basis produknya sudah lebih maju daripada catatan dalam dokumen sumber: repository saat ini memuat empat series dan 20 lesson, halaman bilingual, CTA ke layanan, CTA ke prioritas, serta Website Growth Roadmap. Hambatan utamanya bukan jumlah ide course, melainkan penyelesaian materi, bukti, resource, alur konversi, dan measurement.",
        after=9,
    )
    add_para(
        doc,
        "Draft ini merekomendasikan pilot dua track utama dengan 10 lesson yang sudah ada, dua resource praktis, dan satu alur roadmap bersama. Track lain dipertahankan sebagai backlog terstruktur dan hanya diprioritaskan setelah data pilot menunjukkan kebutuhan. Dengan pendekatan ini, Kultivate menguji apakah Learn menghasilkan pemahaman dan konsultasi yang lebih berkualitas sebelum menambah volume konten.",
        after=10,
    )

    add_heading(doc, "Hasil yang ingin dicapai", level=2)
    add_bullets(doc, [
        "Calon client memahami masalah bisnis dan pilihan intervensi sebelum menghubungi Kultivate.",
        "Kultivate dapat melihat topic, role, dan konteks yang mendorong user dari belajar menuju roadmap atau konsultasi.",
        "Konsultasi dimulai dengan konteks yang lebih jelas tanpa menyembunyikan pengetahuan dasar di balik form.",
        "Tim dapat memproduksi, memperbarui, dan mengukur content tanpa harus membangun LMS penuh.",
    ])

    add_heading(doc, "Dasar adaptasi", level=1, page_break=True)
    add_para(
        doc,
        "Rencana sumber memberi arah strategis yang kuat, khususnya prinsip open learning, progressive conversion, dan batas antara pengetahuan gratis dengan eksekusi berbayar. Adaptasi berikut menyesuaikan arah tersebut dengan kondisi repository saat ini dan menurunkan risiko produksi yang terlalu besar pada tahap awal.",
        after=10,
    )

    add_heading(doc, "Kondisi repository saat ini", level=2)
    add_table(
        doc,
        ["Area", "Yang sudah tersedia", "Implikasi untuk planning"],
        [
            ["Struktur", "4 series dan 20 lesson, masing-masing 5 lesson", "Tidak perlu memulai ulang dari struktur 4 series dan 8 lesson yang tercatat di sumber."],
            ["Bahasa", "Konten metadata ID dan EN", "Workflow penulisan harus menjaga parity bahasa sebelum publikasi."],
            ["Pengalaman", "Learn index, halaman series, lesson navigation, related service CTA", "Fokus perubahan pada isi lesson, resource, status, dan alur sesudah belajar."],
            ["Roadmap", "Quiz 5 pertanyaan dengan rekomendasi SEO atau Web Services", "Gunakan sebagai diagnostic bersama; perlu menerima context dari Learn, bukan dibuat ulang."],
            ["Analytics", "Event Learn, roadmap, dan consultation dasar sudah terdaftar", "Tambahkan event minimum dan parameter context; hindari vocabulary baru yang tumpang tindih."],
            ["Asset", "Video ID, duration, poster, dan thumbnail masih kosong", "Bottleneck awal adalah produksi dan approval, bukan menambah lebih banyak judul."],
        ],
        [1.25, 2.35, 3.4],
    )

    add_heading(doc, "Perubahan dari rencana sumber", level=2)
    add_table(
        doc,
        ["Arah sumber", "Adaptasi yang disarankan", "Alasan"],
        [
            ["6 track dan 36 lesson", "Pilot 2 track dan 10 lesson yang sudah ada", "Memvalidasi kualitas content dan funnel sebelum memperbesar beban produksi."],
            ["3 track P1 termasuk Procurement", "Owner dan Marketing menjadi audience awal; Procurement masuk setelah divalidasi", "Repository dan service journey saat ini lebih dekat ke SEO dan Web Services."],
            ["Membangun diagnostic baru", "Perluas Website Growth Roadmap yang sudah ada", "Mengurangi duplikasi UI, logic, QA, dan analytics."],
            ["Progressive data capture sejak resource", "Utamakan privasi; data personal diminta saat menyimpan hasil, follow-up, atau konsultasi", "Roadmap saat ini berjalan di browser dan tidak menyimpan jawaban."],
            ["Video sebagai bentuk utama", "Naskah dan ringkasan lebih dulu, kemudian video setelah review editorial", "Mempercepat validasi dan mengurangi re-recording."],
            ["Banyak event baru", "Pertahankan event yang ada dan tambah tiga event MVP", "Event taxonomy tetap ringkas dan mudah diuji."],
        ],
        [1.6, 2.45, 2.95],
    )

    add_heading(doc, "Prinsip produk", level=1, page_break=True)
    add_numbered(doc, [
        "Mulai dari keputusan bisnis. Judul dan contoh lesson harus membantu user memilih prioritas, bukan mempelajari tool secara abstrak.",
        "Berikan framework secara utuh. Nilai berbayar berasal dari research kontekstual, prioritas, design, development, QA, dan accountability.",
        "Tidak ada hard gate sebelum belajar. Data personal diminta hanya ketika user menerima hasil yang perlu disimpan, ditindaklanjuti, atau dibahas.",
        "Satu lesson menyelesaikan satu keputusan. Setiap lesson berakhir dengan self-check atau next action yang jelas.",
        "Gunakan bukti yang dapat diverifikasi. Angka, hasil client, dan klaim performa harus memiliki sumber, periode, dan batas interpretasi.",
        "Pertahankan jalur keluar yang ringan. User boleh lanjut belajar, melihat layanan, memakai roadmap, atau berhenti tanpa dipaksa konsultasi.",
    ])

    add_heading(doc, "Arsitektur pilot", level=1)
    add_para(
        doc,
        "Pilot menggunakan series dan lesson ID yang sudah ada agar perubahan route dan data tidak membesar. Nama publik dapat disempurnakan tanpa langsung memecah seluruh model. Dua track utama mendapat content lengkap dan funnel penuh. Dua track spesialis ditahan dari promosi utama sampai materi dan review cadence siap.",
        after=9,
    )
    add_table(
        doc,
        ["Prioritas", "Track", "Isi awal", "Treatment"],
        [
            ["P0", "Demand Search dan Prioritas Halaman", "5 lesson dari Search Demand and Discovery", "Selesaikan script, summary, contoh, resource, dan CTA. Jadikan jalur utama SEO."],
            ["P0", "Keputusan Website dan Conversion", "5 lesson dari Website Decisions and Conversion", "Selesaikan content dan scorecard. Jadikan jalur utama Web Services."],
            ["P1", "Kualitas Website dan Migrasi", "5 lesson yang sudah ada", "Pertahankan sebagai track spesialis. Publikasikan setelah bukti dan checklist teknis lolos review."],
            ["P1", "Technical dan AI Search", "5 lesson yang sudah ada", "Review overlap. Pisahkan technical foundation dari AI Search hanya jika kapasitas update tersedia."],
            ["Gate", "Scope dan Evaluasi Proyek", "Belum ada series khusus", "Mulai sebagai guide dan vendor scorecard. Ubah menjadi course bila audience Procurement terbukti relevan."],
        ],
        [0.65, 1.85, 1.75, 2.75],
    )

    add_heading(doc, "Paket content setiap lesson", level=2)
    add_bullets(doc, [
        "Judul yang menyatakan keputusan atau masalah yang dibahas.",
        "Ringkasan tertulis yang dapat dipahami tanpa menonton video.",
        "Script atau transcript dengan satu framework utama dan satu contoh yang dapat diverifikasi.",
        "Tiga sampai lima pertanyaan self-check untuk membantu user menilai konteksnya sendiri.",
        "Next step yang sesuai: lesson berikutnya, resource, roadmap, atau layanan terkait.",
        "Metadata produksi: status, format, owner, last reviewed, sumber klaim, dan asset approval.",
    ])

    add_heading(doc, "Dua resource pilot", level=2)
    add_table(
        doc,
        ["Resource", "Fungsi", "Bentuk awal", "Jalur berikutnya"],
        [
            ["Search Priority Map", "Mengubah demand dan intent menjadi daftar halaman prioritas", "Worksheet 2 halaman dengan contoh dan scoring sederhana", "Website Growth Roadmap dengan source search demand"],
            ["Website Decision Scorecard", "Membandingkan fix, redesign, atau rebuild berdasarkan constraint", "Scorecard 2 halaman dengan interpretasi hasil", "Website Growth Roadmap dengan source website decision"],
        ],
        [1.55, 2.0, 2.05, 1.4],
    )

    add_heading(doc, "Alur funnel pilot", level=1, page_break=True)
    add_para(
        doc,
        "Funnel pilot memanfaatkan komponen yang sudah ada. Perubahan terpenting adalah membawa context course dan lesson sampai ke roadmap serta consultation form, sehingga Kultivate dapat mengetahui perjalanan user tanpa meminta data personal di awal.",
        after=9,
    )
    add_table(
        doc,
        ["Tahap", "Nilai untuk user", "Data minimum", "CTA utama"],
        [
            ["Learn index", "Memilih masalah yang paling relevan", "Locale dan series click", "Buka track"],
            ["Lesson", "Framework, contoh, dan self-check", "Series, lesson, source, completion", "Lanjut atau gunakan resource"],
            ["Resource", "Worksheet atau scorecard yang dapat dipakai sendiri", "Resource click atau download", "Buka roadmap dengan context"],
            ["Roadmap", "Rekomendasi awal berdasarkan jawaban", "Jawaban tetap di browser pada MVP; simpan source dan hasil secara anonim", "Lihat layanan atau lanjut konsultasi"],
            ["Consultation", "Pembahasan prioritas berdasarkan konteks nyata", "Nama, contact, URL, problem, timeline, source, recommended service", "Kirim permintaan konsultasi"],
        ],
        [1.0, 1.9, 2.15, 1.95],
    )

    add_heading(doc, "Aturan CTA", level=2)
    add_bullets(doc, [
        "Lesson awal memprioritaskan lanjut belajar; consultation CTA tidak perlu muncul berulang di setiap blok.",
        "Resource CTA muncul setelah framework dapat langsung diterapkan.",
        "Roadmap CTA muncul ketika user sudah memahami problem dan mampu memberi input yang bermakna.",
        "Consultation CTA muncul pada hasil roadmap, akhir track, dan kondisi intent tinggi.",
        "Related service CTA tetap tersedia sebagai jalur sekunder untuk user yang ingin membaca scope layanan terlebih dahulu.",
    ])

    add_heading(doc, "Data model minimum", level=2)
    add_table(
        doc,
        ["Entitas", "Field yang perlu ditambah", "Tujuan"],
        [
            ["Series", "Audience, businessProblem, status, resource, roadmapContext, relatedService, lastReviewed", "Mengelola segmentasi, CTA, dan lifecycle tanpa hardcoding per halaman."],
            ["Lesson", "Format, contentStatus, transcript, sources, selfCheck, nextStep, lastReviewed", "Memisahkan metadata editorial dari asset video dan memastikan setiap lesson siap dipublikasikan."],
            ["Resource", "ID, locale, file or route, relatedSeries, consentMode, version", "Menjaga asset dapat dilacak dan diperbarui tanpa menyalin CTA di banyak komponen."],
            ["Funnel context", "Locale, series, lesson, resource, roadmapResult, source", "Meneruskan context sampai consultation form dan analytics."],
        ],
        [1.05, 3.65, 2.3],
    )

    add_heading(doc, "Measurement", level=1)
    add_lead_para(
        doc,
        "Ukuran utama sementara. ",
        "Jumlah consultation submit yang berkualitas dan memiliki touchpoint Learn, bukan jumlah view atau jam video.",
        after=9,
    )
    add_table(
        doc,
        ["Jenis", "Metric", "Cara membaca"],
        [
            ["Progress", "Series open, lesson start, lesson complete", "Menilai apakah user benar-benar bergerak melewati materi awal."],
            ["Application", "Resource click or download", "Menilai apakah framework cukup berguna untuk diterapkan."],
            ["Intent", "Roadmap start, complete, result", "Menilai perpindahan dari belajar umum menuju diagnosis konteks."],
            ["Conversion", "Consultation click, form start, form submit", "Menilai handoff menuju percakapan komersial."],
            ["Quality", "Kecocokan problem, readiness, scope clarity, next action", "Ditag manual sesudah konsultasi sampai volume cukup untuk sistem yang lebih formal."],
        ],
        [1.0, 2.5, 3.5],
    )
    add_para(
        doc,
        "Event MVP yang perlu ditambahkan hanya learn_lesson_complete, learn_resource_click, dan learn_resource_download. Event yang sudah ada tetap dipakai. Semua event roadmap dan consultation perlu membawa parameter source ketika user datang dari Learn. Target numerik ditetapkan setelah baseline pilot tersedia.",
        after=9,
    )

    add_heading(doc, "Rencana pelaksanaan delapan minggu", level=1, page_break=True)
    add_para(
        doc,
        "Durasi berikut adalah perkiraan kerja untuk direvisi berdasarkan kapasitas owner, content, design, dan development. Pekerjaan content dan product dapat berjalan paralel jika owner berbeda.",
        after=9,
    )
    add_table(
        doc,
        ["Minggu", "Fokus", "Output", "Kriteria selesai"],
        [
            ["1", "Alignment dan inventory", "Scope pilot, audience awal, lesson map, asset inventory, analytics baseline", "Dua track, owner, format, dan release rule disetujui."],
            ["2", "Content standard dan data model", "Lesson template, source policy, metadata baru, status workflow", "Satu lesson contoh tampil lengkap dalam ID dan EN."],
            ["3-4", "Produksi track pilot", "10 naskah dan ringkasan, self-check, contoh, 2 draft resource", "Semua lesson lolos review fakta dan editorial; tidak ada placeholder publik."],
            ["4-5", "Funnel integration", "Resource CTA, context ke roadmap, context ke consultation form", "Source series, lesson, dan result dapat diuji end to end."],
            ["6", "Analytics dan QA", "Event spec, dashboard view, accessibility and responsive QA", "Event tidak double fire; keyboard, mobile, locale, dan reduced motion lolos."],
            ["7", "Soft launch", "Rilis terbatas, feedback log, consultation tagging", "Masalah kritis ditutup dan data dasar terkumpul."],
            ["8", "Review dan keputusan ekspansi", "Laporan pilot, revisi lesson, keputusan backlog", "Keputusan lanjut, revisi, atau hentikan didasarkan pada bukti pilot."],
        ],
        [0.65, 1.25, 2.55, 2.55],
    )

    add_heading(doc, "Pembagian peran", level=2)
    add_table(
        doc,
        ["Peran", "Tanggung jawab"],
        [
            ["Product owner", "Menetapkan audience, prioritas track, consultation capacity, dan keputusan release."],
            ["Subject matter owner", "Memastikan framework, contoh, batas klaim, dan execution boundary akurat."],
            ["Content editor", "Menyusun script, summary, transcript, bahasa ID dan EN, serta source notes."],
            ["Design and development", "Mengembangkan metadata, resource surface, context passing, responsive UI, dan accessibility."],
            ["Analytics and QA", "Mendefinisikan event, menguji funnel, memeriksa data, dan mencatat defect."],
        ],
        [1.65, 5.35],
    )
    add_para(
        doc,
        "Satu orang dapat memegang beberapa peran. Namun approval content dan approval release sebaiknya tetap dibedakan agar materi tidak dipublikasikan hanya karena asset teknis sudah tersedia.",
        after=9,
    )

    add_heading(doc, "Kriteria selesai pilot", level=1)
    add_bullets(doc, [
        "Dua track utama memiliki 10 lesson lengkap dalam ID dan EN atau mengikuti release rule bahasa yang sudah disetujui.",
        "Setiap lesson memiliki summary, framework, contoh yang dapat diverifikasi, self-check, next step, dan last reviewed date.",
        "Dua resource pilot dapat digunakan tanpa consultation dan memiliki version owner.",
        "Learn meneruskan source context ke roadmap dan consultation form.",
        "Roadmap tetap dapat digunakan tanpa mengirim atau menyimpan jawaban personal pada tahap MVP.",
        "Event utama teruji tanpa duplicate firing dan dapat dibaca pada analytics yang dipilih.",
        "Semua halaman lolos mobile, keyboard, screen-reader structure, localization, dan reduced-motion QA.",
        "Tidak ada video kosong, fake case study, placeholder visual publik, atau klaim performa tanpa sumber.",
        "Ada prosedur manual untuk menilai kualitas consultation yang datang dari Learn.",
    ])

    add_heading(doc, "Risiko utama dan mitigasi", level=2)
    add_table(
        doc,
        ["Risiko", "Dampak", "Mitigasi"],
        [
            ["Scope content membesar", "Launch tertunda dan kualitas tidak konsisten", "Batasi pilot pada 10 lesson dan dua resource; gunakan expansion gate."],
            ["Track terlihat lengkap tetapi asset kosong", "User kehilangan kepercayaan", "Gunakan contentStatus dan sembunyikan lesson sampai definition of done terpenuhi."],
            ["Course terlalu generik", "Engagement ada tetapi consultation tidak relevan", "Mulai dari keputusan bisnis, contoh nyata, dan self-check per lesson."],
            ["Data diminta terlalu dini", "Drop-off dan risiko consent", "Pertahankan open learning; minta data personal hanya untuk save, follow-up, atau consultation."],
            ["Analytics tidak dapat ditafsirkan", "Tim mengejar view tanpa memahami intent", "Dokumentasikan event, source context, funnel stage, dan quality tagging."],
            ["AI Search cepat usang", "Materi menyesatkan atau menurunkan credibility", "Pisahkan track, beri last reviewed date, dan tetapkan owner serta review cadence."],
        ],
        [1.75, 2.2, 3.05],
    )

    add_heading(doc, "Daftar untuk revisi Anda", level=1, page_break=True)
    add_para(
        doc,
        "Gunakan daftar ini untuk menandai keputusan, koreksi, dan batas scope. Pilihan di kolom kedua adalah posisi awal draft, bukan keputusan final. Pertanyaan pendalaman akan disusun setelah revisi Anda diterima.",
        after=9,
    )
    add_table(
        doc,
        ["Status", "Keputusan yang perlu ditetapkan", "Pilihan awal draft"],
        [
            ["[ ]", "Tujuan utama Learn", "Meningkatkan kualitas consultation dan membantu buyer memahami prioritas."],
            ["[ ]", "Audience pilot", "Owner atau Founder dan Marketing Lead lebih dulu."],
            ["[ ]", "Track launch", "Search Demand and Discovery serta Website Decisions and Conversion."],
            ["[ ]", "Jumlah lesson launch", "10 lesson yang sudah ada, bukan 18 atau 36."],
            ["[ ]", "Posisi Procurement", "Mulai sebagai panduan dan scorecard; course penuh setelah validasi."],
            ["[ ]", "Urutan bahasa", "Tulis master ID, adaptasi EN, lalu public release setelah parity disetujui."],
            ["[ ]", "Format content", "Naskah lebih dulu, kemudian video setelah naskah dan contoh disetujui."],
            ["[ ]", "Lead capture resource", "Ungated pada pilot; opt-in ditambahkan hanya jika follow-up flow siap."],
            ["[ ]", "Roadmap", "Perluas quiz yang ada dan bawa context dari Learn."],
            ["[ ]", "Channel consultation", "Form utama; WhatsApp tetap jalur tambahan bila capacity and tracking jelas."],
            ["[ ]", "Bukti dan case study", "Gunakan public evidence lebih dulu; client case hanya dengan permission."],
            ["[ ]", "Analytics platform dan owner", "Tetapkan tool, dashboard owner, dan review cadence sebelum soft launch."],
            ["[ ]", "Consultation capacity", "Sesuaikan volume CTA dan follow-up dengan siapa yang benar-benar menangani call."],
            ["[ ]", "Expansion gate", "Tambah track hanya setelah content quality, funnel signal, dan operational capacity memadai."],
        ],
        [0.55, 2.75, 3.7],
    )

    add_heading(doc, "Yang sengaja tidak masuk pilot", level=2)
    add_bullets(doc, [
        "Account system, progress sync lintas device, badges, certificate, dan gamification.",
        "Enam diagnostic terpisah untuk setiap track.",
        "Email nurture otomatis sebelum consent, content owner, dan delivery system siap.",
        "AI audit yang memberi diagnosis teknis atau SEO tanpa data yang dapat diverifikasi.",
        "Rekaman 36 video sebelum format lesson dan conversion path terbukti bekerja.",
        "Perubahan besar pada route dan slug hanya untuk menyesuaikan nama course baru.",
    ])

    add_heading(doc, "Tahap setelah revisi", level=2)
    add_numbered(doc, [
        "Anda mengoreksi daftar keputusan, scope, urutan prioritas, dan asumsi yang tidak sesuai.",
        "Pertanyaan pendalaman disusun hanya untuk area yang belum terjawab atau masih saling bertentangan.",
        "Draft v1.0 dibuat dengan timeline, owner, dependency, acceptance criteria, dan backlog final.",
        "Implementation brief untuk content dan website dipisahkan agar dapat dieksekusi tanpa membawa seluruh dokumen strategi.",
    ])

    add_para(
        doc,
        "Basis adaptasi: dokumen Kultivate Learn Free Course Content and Consultation Funnel Plan tanggal 9 September 2026 dan kondisi repository lokal yang diperiksa pada 10 September 2026.",
        before=14,
        after=0,
    )
    doc.paragraphs[-1].runs[0].font.size = Pt(8.5)
    doc.paragraphs[-1].runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()




