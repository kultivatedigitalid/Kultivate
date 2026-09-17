from pathlib import Path
import ast,json
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(r'C:\Users\Joshua\OneDrive\Documents\Kultivate')
OLD=ROOT/'tmp/c01l01-build'
BUILD=ROOT/'tmp/c01l01-v2'
OUT=ROOT/'outputs/Kultivate_Learn_C01L01_Sample_v2_Bahasa_Santai_2026-09-14'
BUILD.mkdir(parents=True,exist_ok=True)
OUT.mkdir(parents=True,exist_ok=True)
# Reuse the existing document layout helpers, without running the old authoring code.
tree=ast.parse((OLD/'build_documents.py').read_text(encoding='utf8'))
helpers=[n for n in tree.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef))]
exec(compile(ast.Module(body=helpers,type_ignores=[]),str(OLD/'build_documents.py'),'exec'))
DATE='14 September 2026'
SOURCES=json.loads((OLD/'sources.json').read_text(encoding='utf8'))
source_notes=[
 'Update terakhir 10 Desember 2025. Kita pakai buat jelasin fungsi SEO, konten yang berguna, dan kenapa ranking pertama nggak bisa dijanjikan. Dipakai di Bab 1 dan 4.',
 'Halaman ini nggak menampilkan tanggal update. Kita pakai buat bedain impressions sama clicks. Cara hitung impressions bisa beda sesuai jenis hasil pencariannya. Dipakai di Bab 3.',
 'Halaman ini nggak menampilkan tanggal update. Key event mencatat tindakan yang dianggap penting buat bisnis. Soal calon pembelinya cocok atau nggak tetap perlu kita cek sendiri. Dipakai di Bab 3.',
 'Edisi Local Results Netherlands, halaman PDF 3, 5–6. Tanggal terbit nggak tercantum. Risetnya membahas orang yang bisa bolak-balik mencari pilihan dan menilai pilihan itu. Ini riset konsumen, jadi bukan bukti kalo semua pembeli B2B di Indonesia punya pola yang sama. Dipakai di Bab 2.',
 'Kita ambil cara halaman ini nunjukin urutan tahap belajar, resources, dan jalur bantuan yang lebih personal. Ini referensi buat cara ngajarnya, bukan sumber fakta SEO atau bukti hasil bisnis Kultivate.',
 'Kita ambil hubungan antara phase, week, lesson, itinerary, dan worksheet. Halaman ini juga punya jalur ke layanan. Kita nggak meniru gaya bicara, copy, atau personanya.'
]
for s,n in zip(SOURCES,source_notes):s[4]=n

def paras(d,items):
    for item in items:p(d,item)

mod=docbase('MODUL PEMBELAJARAN')
mod.add_heading('Gimana pencarian bisa buka peluang bisnis',0)
p(mod,'How Search Creates Business Opportunities',size=12)
p(mod,'Di lesson ini, kita mau lihat gimana search bisa bantu calon pelanggan sebelum mereka hubungi bisnis kita. Kita bakal bahas apa yang mereka butuhkan, jawaban apa yang perlu kita kasih, dan apa aja yang masih harus dicek sebelum bilang ada peluang bisnis.')
p(mod,'Course: 01 — SEO for Business Growth: From Search to Sales',lead='Course:')
p(mod,'Lesson: 01 dari 06 di Course 01. Satu lesson punya satu video utama, dengan target durasi 6–12 menit.',lead='Lesson:')
p(mod,'Audience: Owner atau Founder, Marketing Lead, dan SEO Lead. Belum pernah belajar SEO juga bisa ikut.',lead='Audience:')
p(mod,'Core decision atau mental model: Kita lihat dulu keputusan apa yang bisa terbantu lewat search, lalu cek bukti berikutnya sebelum bilang pencarian itu jadi peluang bisnis.',lead='Core decision atau mental model:')
p(mod,'Learning objective: Kita paham gimana calon pembeli cari tahu, bandingin pilihan, dan ambil keputusan sebelum inquiry masuk. Kita juga bisa lihat hubungan search visibility dengan kebutuhan yang cocok sama bisnis, tanpa nyamain traffic dengan revenue atau pendapatan.',lead='Learning objective:')
p(mod,'Expected learner outcome: Satu peta awal buat satu layanan. Isinya empat momen pencarian, jawaban atau bukti yang calon pembeli butuhkan, satu hal yang masih kurang jelas, dan langkah buat cek hal itu.',lead='Expected learner outcome:')
p(mod,'Pengajar: Hans Galdino. Modul ini jadi bahan buat pahami materinya lebih dulu. Setelah itu, Hans bisa pakai slide buat jelasin dengan bahasanya sendiri.',lead='Pengajar:')
h(mod,'Posisi kita di Kultivate Learn')
paras(mod,[
 'Kultivate Learn punya 6 Courses. Tiap Course punya 6 Lessons, jadi totalnya 36 Lessons dengan 1 video utama buat tiap lesson. Tiap course ngumpulin masalah bisnis yang saling berkaitan. Kita bisa pilih course yang paling dekat sama masalah kita sekarang, lalu ikuti urutan lesson di dalamnya supaya lebih gampang nyambung ke materi berikutnya.',
 'Satu lesson sengaja fokus ke satu keputusan atau mental model, yaitu cara kita memahami masalah. Tujuannya supaya kita nggak harus cerna terlalu banyak konsep sekaligus. Kita bisa pahami dulu, coba pakai, baru lanjut. Nomor Course dan Lesson juga bantu kita tahu posisi kita lagi di mana. Angka enam ini pilihan Kultivate buat nyusun kurikulum, bukan angka yang secara riset pasti paling bagus.'
])
h(mod,'1 Orang cari sesuatu karena ada kebutuhan')
paras(mod,[
 'Bayangin kita nerima laporan kalo website kita makin sering muncul di hasil pencarian dan jumlah kunjungannya meningkat. Jelas aja itu kabar baik. Tapi buat nilai manfaatnya untuk bisnis, kita juga masih harus paham siapa aja yang datang, apa yang mereka cari, dan apakah website kita bisa bantu mereka ambil keputusan yang berkaitan dengan layanan kita. Nah, makanya laporan traffic aja masih belum bisa jawab pertanyaan itu.',
 'Demand itu kebutuhan atau keinginan orang terhadap suatu solusi. Search demand adalah bagian dari kebutuhan itu yang kelihatan lewat pencarian. Sementara traffic berarti kunjungan ke website. Jadi, yang satu bicara soal kebutuhan orang, yang satu lagi soal aktivitas di website. Kebutuhannya bisa aja udah ada meski orang itu nggak pakai search. Misalnya, dia langsung tanya ke teman kerja buat cari rekomendasi.',
 'Kenapa kita perlu bedain? Karena orang yang baca topik seputar bidang usaha kita belum tentu butuh layanan kita. Misalnya, artikel cara bersihin AC sendiri mungkin dibaca pemilik rumah. Kalo bisnis kita khusus pegang kontrak perawatan AC kantor, kunjungan tadi belum tentu cocok sama layanan kita. Tapi orang yang baru belajar soal perawatan AC kantor juga jangan langsung dianggap nggak berguna. Dia mungkin belum siap beli sekarang, tapi lagi cari bekal buat ambil keputusan nanti.',
 'SEO, singkatan dari search engine optimization, bantu mesin pencari pahami konten kita. SEO juga bantu orang nemuin website dan menilai apakah isinya layak mereka kunjungi. Google jelasin fungsi dasarnya di SEO Starter Guide. [S1] Kita pakai istilah search visibility buat nyebut kehadiran bisnis di hasil pencarian. Di lesson ini, fokus kita adalah pencarian organik, yaitu hasil pencarian yang bukan iklan, dan perannya sebelum orang hubungi bisnis kita.',
 'Ada juga istilah commercial intent. Maksudnya, tujuan pencarian yang berkaitan dengan niat menimbang atau membeli solusi. Misalnya, orang cari penyedia jasa, mau tahu layanan apa aja yang termasuk, atau lagi bandingin penawaran. Itu bisa jadi petunjuk. Tapi dari kata yang dia ketik aja, kita belum tahu apakah kebutuhannya cocok, apakah dia boleh ambil keputusan beli, atau apakah dia udah siap lanjut. Konteksnya tetap harus kita cek.',
 'Jadi sebelum ngejar lebih banyak kunjungan, kita perlu bisa jawab satu hal: pencarian ini nyambung nggak sama kebutuhan yang bisa bisnis kita bantu? Kita belum perlu hitung besarnya pasar atau susun daftar keyword yang mau dikerjain. Untuk lesson pertama ini, kita pahami dulu hubungannya.'
])
h(mod,'2 Empat momen sebelum orang hubungi kita')
p(mod,'Inquiry itu pertanyaan atau permintaan yang beneran masuk ke bisnis kita. Bentuknya bisa pesan soal layanan atau permintaan penawaran. Sebelum sampai ke situ, biasanya ada hal-hal yang calon pembeli masih pengin tahu. Kita pakai empat momen buat melihatnya: Discover, Learn, Compare/Validate, dan Act. Ini Kultivate framework, cara kita bantu susun pertanyaan. Ini bukan pembagian resmi dari Google.')
table(mod,['Momen','Yang lagi dicari calon pembeli','Yang bisa kita bantu'],[
 ('Discover','Masalahnya apa dan solusi apa yang mungkin ada.','Bantu orang kenali masalah dan hubungannya sama layanan kita.'),
 ('Learn','Solusinya kerja gimana dan apa aja yang perlu dipikirin.','Jelasin pilihan dengan bahasa yang gampang dipahami.'),
 ('Compare/Validate','Pilihan mana yang cocok dan bisa dipercaya.','Jelasin cakupan, batas layanan, dan bukti yang bisa dicek.'),
 ('Act','Gimana cara mulai hubungi bisnis.','Bikin langkah kontak dan proses setelahnya jadi jelas.')
],[3.0,6.5,7.3])
paras(mod,[
 'Di Discover, orang bisa sadar ada masalah yang perlu diberesin, tapi belum tahu nama solusi atau siapa yang bisa bantu. Pertanyaannya masih dasar. Di Learn, dia mulai cari tahu pilihan yang ada. Kita bantu dia paham bedanya, termasuk kapan sebuah layanan belum perlu dipakai. Jadi, nggak semua konten awal harus langsung minta orang hubungi sales.',
 'Di Compare/Validate, calon pembeli mulai cek apakah pilihannya cocok dan bisa dipercaya. Compare artinya bandingin pilihan. Validate artinya cek lagi apakah alasan buat percaya itu cukup. Misalnya, dia dapat nama vendor dari teman, lalu cari lewat Google buat lihat layanan atau bukti kerjanya. Di sini search bantu dia cek pilihan, walaupun awalnya dia tahu vendor itu dari rekomendasi.',
 'Di Act, yang dia pikirin lebih ke langkah berikutnya. Harus siapin info apa? Siapa yang bakal balas? Perlu buat janji dulu atau cukup kirim kebutuhan? Buat bisnis jasa, langkah ini bisa berupa inquiry. Tapi orang yang udah hubungi kita belum tentu bakal tanda tangan kontrak.',
 'Empat momen ini nggak harus dilewati satu per satu secara urut. Orang bisa langsung cari nama vendor, balik lagi buat pelajari pilihannya, lalu tanya pendapat orang lain. Riset Think with Google juga bahas orang yang bisa bolak-balik mencari pilihan dan menilai pilihan itu. [S4] Jadi, gambar alur yang lurus jangan dibaca terlalu kaku. Kita pakai gagasan itu buat bantu memahami contoh B2B di sini. Bukan berarti riset tadi membuktikan semua pembeli B2B punya pola yang sama.',
 'Dari empat momen ini, coba lihat pertanyaan mana yang belum kita jawab dengan baik. Tapi jangan langsung anggap kita harus bikin empat halaman baru. Satu halaman bisa bantu lebih dari satu momen. Soal jenis halaman dan susunannya, nanti kita bahas di lesson berikutnya.'
])
h(mod,'3 Traffic sampai peluang bisnis perlu dicek satu per satu')
paras(mod,[
 'Tadi kita lihat apa yang lagi dipikirin calon pembeli. Sekarang kita lihat buktinya dari sisi bisnis. Website muncul di search, orang datang, inquiry masuk, lalu ada peluang yang cocok. Kedengarannya nyambung, tapi tiap bagian punya arti yang beda. Kita nggak bisa pakai satu angka buat mewakili semuanya.',
 'Qualified opportunity itu peluang yang udah kita cek kecocokannya dan layak dibahas lebih lanjut. Misalnya, calon pembeli butuh layanan yang memang kita punya, lokasinya bisa kita layani, dan ada langkah berikutnya yang masuk akal. Bisnis kita perlu sepakat sendiri soal kriteria cocoknya. Kalo orang cuma tanya layanan yang nggak kita punya, inquiry memang masuk, tapi belum jadi peluang yang cocok. Bahkan kalo udah cocok pun, belum tentu berujung penjualan.'
])
table(mod,['Sinyal','Yang udah kita tahu','Yang masih perlu dicek'],[
 ('Visibility','Bisnis atau kontennya muncul di search. Impressions bisa jadi salah satu tandanya.','Orang itu memang butuh layanan kita atau nggak.'),
 ('Kunjungan relevan','Kunjungannya kelihatan nyambung sama kebutuhan yang bisa kita bantu.','Orangnya udah mau atau siap beli belum.'),
 ('Inquiry','Ada pesan atau permintaan yang beneran kita terima.','Kebutuhan dan kondisi calon pembelinya cocok atau nggak.'),
 ('Qualified opportunity','Kecocokannya udah dicek dan ada langkah berikutnya yang disepakati.','Peluang ini akhirnya jadi penjualan atau nggak.')
],[3.4,6.3,7.1])
paras(mod,[
 'Di Search Console, impressions dan clicks itu beda. Impressions bicara soal kemunculan, dengan cara hitung yang bisa beda sesuai jenis hasil pencarian. Clicks adalah klik dari Google ke website. [S2] Dari sini kita bisa lihat aktivitas di search. Tapi buat tahu kaitannya sama bisnis, kita juga perlu lihat halaman yang dikunjungi dan pertanyaan yang masuk. Jumlah impressions bukan jumlah calon pembeli unik. Jumlah klik juga nggak otomatis sama dengan jumlah inquiry.',
 'Google Analytics punya istilah key event, yaitu tindakan yang kita anggap penting buat bisnis dan kita catat sebagai event. [S3] Contohnya, orang klik tombol kontak. Itu bisa jadi tanda ketertarikan. Tapi kita belum tahu apakah dia beneran kirim pesan dan pesannya sampai ke bisnis kita. Begitu juga kalo formulir tercatat terkirim. Kebutuhannya tetap harus dicek oleh tim. Label event nggak bisa jawab sendiri apakah calon pembeli ini cocok.',
 'Bayangin ada inquiry masuk, tapi orangnya minta layanan yang nggak kita tawarkan. Traffic ada, inquiry juga ada, tapi kecocokannya belum ada. Di sisi lain, orang bisa nemuin penjelasan kita lewat search, simpan link-nya, lalu balik dan hubungi kita lewat kanal lain. Kalo cuma lihat kunjungan terakhir, kita nggak tahu seluruh ceritanya. Attribution adalah cara kita membagi kredit atau kontribusi ke interaksi pemasaran yang ikut berperan. Model dan cara ukurnya nanti kita bahas di Lesson 06.',
 'Urutan visibility, kunjungan relevan, inquiry, dan qualified opportunity di sini bantu kita tahu apa yang perlu dicek. Diagramnya belum jadi laporan funnel yang ngikutin orang yang sama dari awal sampai akhir. Jadi kalo traffic naik tapi peluangnya belum jelas, kita cari dulu bagian mana yang belum ada buktinya. Jangan buru-buru bilang SEO-nya berhasil atau gagal.'
])
h(mod,'4 Kita coba pakai di satu layanan')
p(mod,'Contoh ini hipotetis. Kita bayangin sebuah bisnis jasa perawatan AC kantor supaya cara mikirnya lebih kebayang. Ini bukan cerita klien Kultivate dan nggak pakai data performa. Anggap bisnis ini melayani perawatan berkala untuk kantor di area tertentu. Pencarian di bawah juga cuma ilustrasi, bukan keyword hasil riset atau data pelanggan.')
table(mod,['Momen','Contoh pencarian hipotetis','Jawaban yang dia butuhkan'],[
 ('Discover','AC kantor sering bermasalah','Gimana memahami masalahnya dan kapan perlu bantuan penyedia jasa.'),
 ('Learn','Perawatan AC kantor berkala termasuk apa aja','Apa aja yang dikerjakan dan info apa yang perlu disiapkan.'),
 ('Compare/Validate','Vendor perawatan AC kantor dan cakupan layanan','Area yang dilayani, pekerjaan yang termasuk, batasnya, dan bukti yang bisa dicek.'),
 ('Act','Permintaan survei perawatan AC kantor','Gimana cara kirim kebutuhan dan apa yang terjadi setelah itu.')
],[3.0,6.2,7.6])
paras(mod,[
 'Di momen Learn, calon pembeli masih perlu paham dulu sebelum bisa bandingin penawaran. Kalo halaman kita isinya cuma ajakan minta harga, pertanyaan dia belum tentu terjawab. Di Compare/Validate, penjelasan umum juga bisa kurang. Dia mungkin pengin tahu apakah lokasi kantornya bisa dilayani dan pekerjaan apa aja yang termasuk. Jadi, jawaban yang kita kasih perlu ngikutin hal yang lagi dia putuskan.',
 'Sekarang kita coba isi satu bagian worksheet. Kita pilih Compare/Validate. Pertanyaan calon pembelinya: apakah vendor ini bisa layani lokasi kantor dan kebutuhan perawatan kami? Berarti dia perlu info area layanan dan batas pekerjaannya. Dalam contoh ini, kita cek halaman dan nemuin tombol kontak, tapi nggak nemuin info area layanan. Nah, itu observasi: hal yang beneran kita lihat di halaman. Kalo kita bilang calon pembeli akhirnya ragu lalu batal hubungi bisnis, itu masih hipotesis. Artinya, dugaan yang perlu kita cek.',
 'Buat cek dugaan tadi, kita bisa lihat pertanyaan yang masuk dan ngobrol sama orang yang pegang inquiry. Apakah calon pembeli memang sering tanya soal lokasi yang dilayani? Kalo iya, ada alasan buat bikin info itu lebih jelas. Kalo nggak, kita perlu cari tahu apakah ada hal lain yang lebih mereka butuhkan. Kita belum bisa bilang tambahan info itu pasti bikin inquiry naik, apalagi hitung nilai penjualannya.',
 'Misalnya setelah itu ada permintaan survei masuk. Tim tetap perlu cek lokasi, jenis kebutuhan, dan apakah ada langkah berikutnya yang bisa dijalankan. Baru dari situ kita bisa nilai peluangnya sesuai kriteria bisnis. Jadi, pencarian yang kelihatan punya niat beli dan tombol kontak yang jelas tetap belum menjamin penjualan.',
 'Sekarang coba pakai di bisnis Anda. Pilih satu layanan utama di Search-to-Sales Mapping Worksheet bagian C01L01, lalu jawab empat pertanyaan tentang Discover, Learn, Compare/Validate, dan Act. Kalo ada pertanyaan yang pernah Anda dengar langsung dari calon pelanggan, pakai itu dulu. Kalo masih perkiraan, tulis hipotesis. Kalo belum tahu, tulis belum tahu. Nggak perlu maksa punya jawaban pasti cuma supaya semua kolom terisi.',
 'Setelah itu, pilih satu pertanyaan yang menurut kita belum terjawab dengan baik. Catat apa yang beneran kita lihat, apa yang masih kita duga, dan bukti apa yang bisa bantu kita cek. Hal yang masih kurang atau belum jelas ini kita sebut gap. Peta awalnya udah berguna kalo kita bisa jelasin satu gap dan tahu apa yang mau dicek berikutnya. Ini belum jadi skor potensi pasar atau alasan buat otomatis bikin konten baru.',
 'Ada saatnya search juga belum jadi hal pertama yang perlu kita kerjain. Kalo kita belum tahu apakah calon pembeli pakai search buat masalah ini, cari tahu dulu kebiasaan mereka. Kalo inquiry yang cocok udah masuk tapi lama dibalas, kita perlu cek proses balas dan tindak lanjutnya sebelum nambah traffic. Google sendiri nggak menjanjikan ranking pertama atau hasil yang sama dari setiap perubahan website. [S1]',
 'Sampai di sini, kita udah punya cara buat bikin peta awal dan cek pertanyaannya. Kalo mau hitung ukuran pasar, susun kumpulan data keyword, riset kompetitor, bikin model attribution, atau tentukan peluang mana yang lebih dulu dikerjakan, kita butuh riset yang khusus buat bisnis itu. Di situlah batas pekerjaan layanan untuk lesson ini. Anda bisa lanjut bareng tim internal atau penyedia jasa sesuai kebutuhan. Nggak harus pakai Kultivate buat memahami dan mencoba cara dasarnya.'
])
p(mod,'Inti yang kita bawa: Lihat keputusan apa yang bisa terbantu lewat search, lalu cek buktinya sampai ke inquiry yang cocok sama bisnis kita. Traffic aja belum cukup buat bilang ada peluang bisnis.',lead='Inti yang kita bawa:')
p(mod,'Setelah ini, kita lanjut ke Lesson 02, Map Customer Demand Before Choosing Keywords, atau Petakan Demand Pelanggan Sebelum Memilih Keyword. Di sana kita mulai susun kebutuhan dan pertanyaan pelanggan dengan lebih rapi. Kalo Anda isi worksheet C01L01, bawa catatannya. Kalo belum, tetap bisa lanjut belajar.')
h(mod,'Sumber dan catatan tambahan')
p(mod,'Sumber web ini dicek pada 14 September 2026 saat sample awal disusun. Revisi v2 fokus ke bahasa, tanpa menambah klaim teknis. Penanda S merujuk ke fakta atau batas penjelasan dari sumber. Framework, pertanyaan, dan cara pakai contohnya adalah hasil pemikiran Kultivate. T1 dan T2 kita pakai sebagai referensi cara mengajar.')
for code,org,title,url,note in SOURCES:
    link(mod,f'{code}  {org} — {title}',url);p(mod,note,size=9)
p(mod,'Fondasi internalnya tetap Free Course Content and Consultation Funnel Plan (9 September 2026), Content Planning for 6 Courses / 36 Lessons (11 September 2026), Content Production Master Prompt (11 September 2026), dan Trust Explanation Framework (11 September 2026). Formatnya ikut brief revisi, dan bahasanya ikut arahan terbaru: santai seperti dosen lagi jelasin ke mahasiswa. Topik, urutan lesson, tujuan resource, dan batas layanan tetap sama.',size=9)

ws=docbase('WORKSHEET')
ws.add_heading('Search to Sales Mapping Worksheet',0)
p(ws,'Bagian C01L01    Gimana pencarian bisa buka peluang bisnis',size=12)
p(ws,'Kita pakai worksheet ini buat bikin peta awal satu layanan dan pilih satu hal yang mau dicek. Kalo ada catatan pelanggan atau isi website yang bisa dipakai, mulai dari situ. Kalo masih dugaan, tulis hipotesis. Kalo belum tahu, tulis belum tahu. Nggak perlu tool berbayar. Sisihkan sekitar 10–15 menit buat coba isi.')
p(ws,'Layanan yang mau kita lihat dan siapa yang kita layani:');lines(ws,1)
p(ws,'Peran Anda dan URL website kalo ada:');lines(ws,1)
p(ws,'Tujuan bisnis atau masalah di pencarian organik yang lagi kita hadapi:');lines(ws,1)
p(ws,'Timeline kalo udah ada: ____________________  Tanggal catatan: ____________________',size=9)
h(ws,'Empat pertanyaan calon pembeli')
for stage,q in [
 ('Discover','Sebelum kenal layanan atau nama bisnis kita, mereka mungkin cari apa?'),
 ('Learn','Apa yang perlu mereka pahami soal solusi sebelum tahu cocok atau nggak?'),
 ('Compare/Validate','Apa yang mereka bandingin, dan bukti apa yang bikin mereka lebih yakin?'),
 ('Act','Apa yang perlu mereka tahu supaya bisa kirim kebutuhan atau hubungi kita?')]:
    p(ws,stage+' — '+q,lead=stage);lines(ws,2)
p(ws,'Di tiap jawaban, tulis mana yang kita tahu dari observasi, mana yang masih hipotesis, dan mana yang belum tahu. Empat momen ini bisa berulang. Kita juga nggak harus bikin empat halaman baru.',size=9)
ws.add_page_break()
h(ws,'Satu hal yang masih perlu kita cek')
p(ws,'Pilih satu momen dari halaman pertama: ____________________________________________')
p(ws,'Pertanyaan calon pembeli yang mau kita bantu jawab:');lines(ws,2)
p(ws,'Jawaban atau bukti yang perlu dia dapat:');lines(ws,2)
p(ws,'Apa yang beneran kita lihat sekarang? Catat juga sumbernya, misalnya halaman website atau catatan inquiry, dan kapan kita cek.');lines(ws,2)
p(ws,'Apa yang masih kita duga atau belum tahu? Pisahkan dari hal yang udah kita lihat tadi.');lines(ws,2)
p(ws,'Bukti apa yang mau kita cari berikutnya, dan siapa yang bisa bantu?');lines(ws,2)
p(ws,'Setelah dicek, apa langkah kita dan kapan mau kita lihat lagi?');lines(ws,2)
h(ws,'Gimana baca hasilnya')
p(ws,'Kalo jawabannya masih dugaan tim sendiri, cek dulu ke orang yang sering ngobrol sama pelanggan atau lihat catatan inquiry. Kalo infonya udah ada di website, cek apakah gampang ditemukan dan memang menjawab kebutuhan mereka. Kalo buktinya udah mendukung adanya gap, catat perbaikan yang masuk akal buat dibahas lebih lanjut. Hasil worksheet ini nggak otomatis berarti kita perlu halaman baru atau anggaran SEO yang lebih besar.')
p(ws,'Kelengkapan catatan: ____ dari 4 momen udah punya catatan. Angka ini cuma nunjukin kelengkapan isian, bukan skor kesehatan SEO atau peluang penjualan.',lead='Kelengkapan catatan:',size=9)
p(ws,'Yang kita bawa setelah selesai: satu layanan, empat momen, satu gap atau hal yang belum tahu, dan satu langkah buat cek. Simpan buat Lesson 02. Soal ukuran pasar, data keyword, kompetitor, attribution, dan urutan peluang tetap butuh riset khusus bisnis.',size=9)

SLIDES=[
 (1,'Gimana pencarian bisa buka peluang bisnis','Bantu kita tahu posisi lesson dan apa yang mau dipelajari.','Judul besar dan struktur seri. Satu lesson fokus ke satu keputusan atau cara memahami masalah.','Metadata dan Posisi kita di Kultivate Learn',[]),
 (2,'Traffic naik, peluangnya belum jelas','Bantu kita lihat kenapa kunjungan yang naik masih perlu dicek hubungannya sama bisnis.','Dua sisi: yang udah kelihatan dan yang masih perlu dicek. Tanpa angka performa.','Bab 1',[]),
 (3,'Demand dan traffic','Bedain kebutuhan orang dengan kunjungan ke website.','Dua definisi besar dan satu pertanyaan soal kecocokan kebutuhan.','Bab 1',[]),
 (4,'Empat momen pencarian','Kenalin Discover, Learn, Compare/Validate, dan Act tanpa bikin urutannya terasa wajib.','Empat simpul dengan panah maju dan panah balik. Label Kultivate framework.','Bab 2',['S4']),
 (5,'Jawaban yang calon pembeli butuhkan','Bantu kita lihat pertanyaan yang beda di tiap momen.','Tabel native empat baris. Momen dan pertanyaannya bisa langsung dibandingkan.','Bab 2',[]),
 (6,'Dari muncul di search sampai peluang yang cocok','Bedain kemunculan, kunjungan relevan, inquiry, dan peluang yang udah dicek.','Alur empat bagian, masing-masing dengan hal yang perlu kita cek. Ini bukan funnel terukur.','Bab 3',['S2']),
 (7,'Klik kontak belum tentu jadi inquiry','Bantu kita bedain tindakan di website dengan pesan yang beneran masuk.','Tiga tindakan berurutan dan satu catatan soal key event.','Bab 3',['S3']),
 (8,'Contoh perawatan AC kantor','Bikin Discover dan Learn lebih kebayang lewat satu contoh bisnis jasa.','Dua momen awal, contoh pencarian, lalu info yang dibutuhkan. Label hipotetis terlihat.','Bab 4',[]),
 (9,'Inquiry masuk, cocoknya tetap perlu dicek','Lanjutin contoh ke Compare/Validate dan Act.','Dua momen berikutnya dan catatan kalo permintaan masuk masih perlu dicek.','Bab 4',[]),
 (10,'Yang kita lihat dan yang masih kita duga','Bantu kita bedain observasi, hipotesis, dan langkah cek berikutnya.','Tiga baris yang masing-masing punya label jelas.','Bab 4',[]),
 (11,'Sekarang coba di layanan Anda','Sambungin contoh ke worksheet dan ke keputusan sederhana yang bisa kita ambil.','Langkah worksheet di kiri, dua kondisi yang perlu dicek dulu di kanan.','Bab 4 dan Worksheet',[]),
 (12,'Search bisa bantu orang ambil keputusan','Tutup dengan satu ide utama dan ajakan lanjut belajar.','Takeaway besar, lalu Lesson 02 dan catatan kalo worksheet boleh dicoba.','Bab 4 bagian penutup',[])
]

plan=docbase('SLIDE PLANNING DAN QUALITY CHECK')
plan.add_heading('Rencana slide dan catatan sample v2',0)
p(plan,'Course 01 Lesson 01    Gimana pencarian bisa buka peluang bisnis',size=12)
p(plan,'Sample v2 ini pakai bahasa yang lebih santai sesuai arahan terbaru. Fokusnya tetap C01L01. Lesson 02–36 belum dikerjakan dan masih nunggu persetujuan sample.')
h(plan,'Acuan bahasa yang kita pakai')
p(plan,'Bahasanya seperti dosen yang lagi jelasin ke mahasiswa: santai, jelas, dan pakai kata sehari-hari. Kita pakai bentuk seperti bayangin, kalo, buat, aja, udah, paham, bantu, dan cek saat memang enak dibaca. Kita nggak perlu maksa semua imbuhan hilang. Istilah teknis tetap dipakai kalau membantu, lalu artinya dijelasin dengan kata yang sederhana.')
p(plan,'Susunan kalimat juga dibuat lebih ngalir. Misalnya, kita mulai dari situasi yang gampang dibayangin, jelasin kenapa itu penting, lalu bantu pembaca lihat apa yang masih perlu dicek. Modul tetap jadi bahan buat pahami konsep. Bahasanya santai, tapi nggak berubah jadi dialog atau naskah yang harus dibaca kata per kata.')
p(plan,'Empat chapter, learning objective, contoh hipotetis AC kantor, dan tujuan Search-to-Sales Mapping Worksheet tetap sama. Definisi teknis, label sumber, serta batas antara pengetahuan umum dan pekerjaan khusus bisnis juga tetap dijaga. Contoh dan sumber tidak ditambah pada revisi bahasa ini.')
p(plan,'Slide ikut pakai bahasa yang lebih ringan. Teksnya tetap singkat karena penjelasan lengkap ada di modul. Deck tetap punya 12 slide dengan diagram, tabel, dan teks yang bisa diedit. Warna Dark Editorial, Arial, dan logo asli Kultivate tetap dipakai.')
h(plan,'Rencana tiap slide')
for n,title,purpose,visual,chapter,codes in SLIDES:
    plan.add_heading(f'Slide {n:02d} {title}',2)
    p(plan,'Purpose: '+purpose,lead='Purpose:',size=10)
    p(plan,f'On-slide content: SLIDE_CONTENT_{n:02d}',lead='On-slide content:',size=10)
    p(plan,'Suggested visual: '+visual,lead='Suggested visual:',size=10)
    if n in [8,9,10]: source='Contoh hipotetis Kultivate. Bukan data pelanggan atau bukti hasil bisnis.'
    elif n==7: source='S3 buat definisi key event. Contoh klik kontak adalah ilustrasi Kultivate.'
    elif n==4: source='Framework dari kurikulum Kultivate. S4 membantu jelasin kenapa perjalanan orang nggak selalu lurus.'
    elif n==6: source='S2 buat impressions dan clicks. Alur cek bukti adalah Kultivate framework.'
    elif n==1: source='Kurikulum dan brief revisi Kultivate. Logo dari asset asli pengguna.'
    else: source='Framework dan penjelasan Kultivate, sesuai kurikulum C01L01.'
    p(plan,'Source: '+source,lead='Source:',size=9)
    p(plan,'Bagian modul: '+chapter,lead='Bagian modul:',size=9)
h(plan,'Sumber dan batas penjelasannya')
p(plan,'Sumber dicek pada 14 September 2026 saat sample awal dibuat. V2 ini revisi bahasa, tanpa riset atau klaim teknis tambahan. Nama sumber dan URL tetap sama supaya gampang dicek lagi.')
for code,org,title,url,note in SOURCES:
    link(plan,f'{code}  {org} — {title}',url);p(plan,note,size=9)
p(plan,'Empat momen pencarian dan alur cek bukti adalah Kultivate framework. Itu bukan model resmi Google atau bukti hubungan sebab-akibat. Semua pencarian, isi halaman, dan inquiry dalam contoh AC kantor bersifat hipotetis. Dugaan soal calon pembeli yang ragu tetap dipisah dari hal yang benar-benar terlihat. Struktur 6 × 6 adalah pilihan kurikulum, bukan angka optimal dari riset.')
h(plan,'Quality check sample v2')
p(plan,'Isi: Fungsi SEO, demand, traffic, commercial intent, inquiry, qualified opportunity, key event, dan attribution tetap punya makna yang sama. Materi berhenti pada peta awal dan langkah cek. Riset keyword, ukuran pasar, audit rinci, dan prioritas peluang tetap butuh konteks bisnis.')
p(plan,'Bahasa: Contoh paragraf dari pengguna jadi patokan. Kalimat dibuat lebih sehari-hari dengan sudut pandang kita. Imbuhan dipakai seperlunya. Istilah teknis tetap dijelasin. Tidak ada dialog, cue produksi, atau naskah kata per kata.')
p(plan,'Struktur dan resource: Empat chapter tetap punya alur yang nyambung. Worksheet tetap dua halaman untuk satu layanan, empat momen, satu gap, dan satu pemeriksaan berikutnya. Kolom observasi, dugaan, dan hal yang belum tahu tetap ada. Angka kelengkapan isian bukan skor SEO atau potensi penjualan.')
p(plan,'Kepercayaan: Contoh hipotetis tetap dilabeli. Fakta sumber, interpretasi, dan cara berpikir Kultivate tetap dibedain. Pengetahuan dasarnya tetap lengkap. Langkah berikutnya adalah Lesson 02, sementara worksheet boleh dicoba tanpa harus pakai layanan Kultivate.')
p(plan,'File dan visual: VALIDATION_RESULT')
h(plan,'Setelah sample ditinjau')
p(plan,'Gaya bahasa di v2 ini jadi acuan untuk materi berikutnya setelah sample disetujui. Struktur chapter dan visual tetap bisa menyesuaikan topiknya. Untuk sekarang, pekerjaan berhenti di C01L01.')

def finish(d,filename):
    for root in [d.styles._element,d._element]:
        for el in list(root.iter(qn('w:pBdr'))):el.getparent().remove(el)
    for pp in d.paragraphs:
        if pp._p.find(qn('w:hyperlink')) is not None:pp.paragraph_format.keep_with_next=True
        if pp.text=='Sumber dan catatan tambahan':pp.paragraph_format.page_break_before=True
        if pp.text.startswith(('Slide ','Purpose:','On-slide content:','Suggested visual:','Source:')):pp.paragraph_format.keep_with_next=True
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                if c.text=='Compare/Validate':
                    c.text='Compare/\nValidate'
                    for r in c.paragraphs[0].runs:r.font.name='Arial';r.font.size=Pt(9.5)
    for sec in d.sections:
        for pp in sec.header.paragraphs:pp.clear()
        for pp in sec.footer.paragraphs:pp.clear()
        pp=sec.footer.paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;pp.paragraph_format.tab_stops.clear_all()
        pp.add_run('Kultivate Learn    ')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');pp._p.append(fld)
        for r in pp.runs:r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
    pp=d.paragraphs[0].insert_paragraph_before('KULTIVATE LEARN   /   C01L01   /   SAMPLE V2 BAHASA SANTAI   /   14 SEPTEMBER 2026')
    pp.paragraph_format.space_after=Pt(12)
    for r in pp.runs:r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
    d.save(OUT/filename)

finish(mod,'C01L01_Modul_Pembelajaran_v2.docx')
finish(ws,'C01L01_Worksheet_Search_to_Sales_v2.docx')
finish(plan,'C01L01_Slide_Planning_dan_Quality_Check_v2.docx')
notes=[]
for n,title,purpose,visual,chapter,codes in SLIDES:
    notes.append({'n':n,'title':title,'purpose':purpose,'visual':visual,'chapter':chapter,'codes':codes,'source':('Contoh hipotetis Kultivate.' if n in [8,9,10] else 'Framework dan penjelasan Kultivate. '+', '.join(codes))})
(BUILD/'slides.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2),encoding='utf8')
(BUILD/'sources.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2),encoding='utf8')
print('Created',OUT)
