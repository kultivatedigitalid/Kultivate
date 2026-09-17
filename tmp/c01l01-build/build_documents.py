from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT=Path(r'C:\Users\Joshua\OneDrive\Documents\Kultivate')
BUILD=ROOT/'tmp/c01l01-build'
OUT=ROOT/'outputs/Kultivate_Learn_C01L01_Sample_2026-09-14'
OUT.mkdir(parents=True,exist_ok=True)
DATE='14 September 2026'
SOURCES=[
 ('S1','Google Search Central','SEO Starter Guide','https://developers.google.com/search/docs/fundamentals/seo-starter-guide','Diperbarui 10 Desember 2025. Definisi SEO, kegunaan konten, dan batas jaminan ranking. Dipakai di Bab 1 dan 4.'),
 ('S2','Google Search Console Help','What are impressions position and clicks','https://support.google.com/webmasters/answer/7042828?hl=en','Tanggal pembaruan tidak ditampilkan. Definisi impressions dan clicks bergantung pada jenis hasil. Dipakai di Bab 3.'),
 ('S3','Google Analytics Help','About key events','https://support.google.com/analytics/answer/9267568?hl=en','Tanggal pembaruan tidak ditampilkan. Key event merekam tindakan yang dianggap penting bagi bisnis. Label ini tidak otomatis menilai kelayakan calon pembeli. Dipakai di Bab 3.'),
 ('S4','Think with Google','Decoding Decisions Making sense of the messy middle','https://www.thinkwithgoogle.com/_qs/documents/11071/DecodingDecisions_NL.pdf','Edisi Local Results Netherlands, halaman PDF 3, 5–6. Tanggal publikasi tidak tercantum. Menjelaskan eksplorasi dan evaluasi yang berulang. Penelitian konsumen ini bukan validasi langsung atas seluruh proses pembelian B2B Indonesia. Dipakai sebagai batas model di Bab 2.'),
 ('T1','Acquisition.com','Improvise','https://www.acquisition.com/training/improvise','Struktur halaman menunjukkan progression per tahap, resources, dan jalur personalized roadmap. Rujukan struktur pengajaran, bukan sumber fakta SEO atau klaim efektivitas Kultivate.'),
 ('T2','Neil Patel','SEO Unlocked','https://neilpatel.com/training/seo-unlocked/','Halaman menampilkan phase, week, lesson, itinerary, worksheet, serta layanan. Rujukan hubungan kurikulum dan aplikasi, tanpa meniru copy atau persona.')
]

def font(run,size=None,bold=None,color='000000'):
    run.font.name='Arial'
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    return run

def docbase(label):
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.85); sec.bottom_margin=Cm(1.85)
    sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.1)
    sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)
    for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s=d.styles[n]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.widow_control=True
    n=d.styles['Normal']; n.font.size=Pt(10.5)
    n.paragraph_format.line_spacing=1.16; n.paragraph_format.space_after=Pt(7)
    d.styles['Title'].font.size=Pt(26); d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(10)
    d.styles['Heading 1'].font.size=Pt(17); d.styles['Heading 1'].font.bold=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(17)
    d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    d.styles['Heading 2'].font.size=Pt(12); d.styles['Heading 2'].font.bold=True
    d.styles['Heading 2'].paragraph_format.space_before=Pt(10)
    h=sec.header.paragraphs[0]; font(h.add_run('KULTIVATE LEARN    /    C01L01    /    '+label),8)
    p=sec.footer.paragraphs[0]; font(p.add_run('Sample untuk persetujuan    •    '+DATE+' '),8)
    p.add_run('\t'); p.paragraph_format.tab_stops.add_tab_stop(Cm(15.8),WD_ALIGN_PARAGRAPH.RIGHT)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    d.core_properties.author='Kultivate'; d.core_properties.subject='C01L01 Sample untuk persetujuan'
    return d

def p(d,text,lead=None,size=None):
    x=d.add_paragraph()
    if lead and text.startswith(lead):
        font(x.add_run(lead),size,bold=True); font(x.add_run(text[len(lead):]),size)
    else: font(x.add_run(text),size)
    return x

def h(d,text): return d.add_heading(text,1)

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    pr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for row_i,vals in enumerate([headers]+rows):
        row=t.rows[0] if row_i==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr(); keep=OxmlElement('w:cantSplit'); trpr.append(keep)
        if row_i==0: trpr.append(OxmlElement('w:tblHeader'))
        for c,txt,w in zip(row.cells,vals,widths):
            c.width=Cm(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'273244' if row_i==0 else ('F3F5F7' if row_i%2 else 'FFFFFF')); cp.append(sh)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'95'); e.set(qn('w:type'),'dxa'); mar.append(e)
            cp.append(mar)
            pp=c.paragraphs[0]; pp.paragraph_format.space_after=Pt(1); pp.paragraph_format.line_spacing=1.1
            font(pp.add_run(txt),9.5,row_i==0,'FFFFFF' if row_i==0 else '000000')
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def link(d,label,url):
    pp=d.add_paragraph(); pp.paragraph_format.space_after=Pt(3)
    rel=pp.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); col=OxmlElement('w:color'); col.set(qn('w:val'),'174F86'); rp.append(col)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'20'); rp.append(sz); r.append(rp)
    tx=OxmlElement('w:t'); tx.text=label; r.append(tx); a.append(r); pp._p.append(a)

def lines(d,count=2):
    for _ in range(count):
        pp=d.add_paragraph('________________________________________________________________________________')
        pp.paragraph_format.space_after=Pt(4); font(pp.runs[0],9,color='A0A0A0')

mod=docbase('MODUL PEMBELAJARAN')
mod.add_heading('Menghubungkan pencarian dengan peluang bisnis',0)
p(mod,'How Search Creates Business Opportunities',size=12)
p(mod,'Modul ini membantu kita menilai peran pencarian sebelum calon pelanggan menghubungi bisnis. Sesudah mempelajarinya, kita dapat menjelaskan kebutuhan yang dilayani oleh sebuah pencarian, dukungan informasi yang diperlukan, dan bukti yang masih perlu diperiksa sebelum menyebutnya peluang bisnis.')
p(mod,'Course: 01 — SEO for Business Growth: From Search to Sales',lead='Course:')
p(mod,'Lesson: 01 dari 06 dalam Course 01. Satu lesson memiliki satu video utama dengan sasaran durasi 6–12 menit.',lead='Lesson:')
p(mod,'Audience: Owner atau Founder, Marketing Lead, dan SEO Lead. Materi dapat diikuti tanpa pengetahuan SEO sebelumnya.',lead='Audience:')
p(mod,'Core decision atau mental model: Nilai peran search berdasarkan keputusan calon pembeli yang dibantunya dan bukti lanjutannya, sebelum menyimpulkan bahwa visibility menjadi peluang bisnis.',lead='Core decision atau mental model:')
p(mod,'Learning objective: Memahami bagaimana discovery, evaluasi, dan keputusan terjadi sebelum inquiry, serta menghubungkan search visibility dengan kebutuhan calon pembeli yang sesuai tanpa menganggap traffic sebagai revenue.',lead='Learning objective:')
p(mod,'Expected learner outcome: Satu peta awal untuk satu layanan, berisi empat momen pencarian, informasi atau bukti yang diperlukan calon pembeli, satu gap yang perlu diperiksa, dan satu tindakan untuk mencari bukti.',lead='Expected learner outcome:')
p(mod,'Pengajar: Hans Galdino. Modul dipelajari sebagai bahan ajar; slide mendukung penjelasan dengan bahasa pengajar sendiri.',lead='Pengajar:')
h(mod,'Posisi kita di Kultivate Learn')
p(mod,'Kultivate Learn terdiri dari 6 Courses. Setiap Course memiliki 6 Lessons, sehingga totalnya 36 Lessons dengan 1 video utama per lesson. Setiap course mengelompokkan masalah bisnis yang saling berhubungan. Kita dapat memilih course sesuai masalah yang sedang dihadapi, lalu mengikuti urutan lesson di dalamnya agar fondasinya terbentuk sebelum masuk ke keputusan berikutnya.')
p(mod,'Satu lesson sengaja berfokus pada satu keputusan atau mental model, yaitu cara memahami sebuah masalah. Dengan begitu, kita tidak perlu mempelajari terlalu banyak konsep sekaligus dan dapat menerapkannya satu per satu. Posisi Course dan Lesson membantu kita mengetahui apa yang sudah dipelajari dan apa langkah berikutnya. Angka enam merupakan pilihan desain kurikulum Kultivate, bukan angka yang diklaim paling optimal secara ilmiah.')

h(mod,'1 Pencarian berawal dari kebutuhan')
p(mod,'Bayangkan kita menerima laporan bahwa website semakin sering muncul di hasil pencarian dan jumlah kunjungannya meningkat. Itu kabar yang berguna. Tapi untuk menilai manfaatnya bagi bisnis, kita masih perlu memahami siapa yang datang, apa yang sedang mereka cari, dan apakah website membantu mereka mengambil keputusan yang berkaitan dengan layanan kita. Laporan traffic saja belum menjawab pertanyaan tersebut.')
p(mod,'Demand adalah kebutuhan atau keinginan terhadap suatu solusi. Search demand merupakan bagian kebutuhan itu yang terlihat melalui pencarian. Traffic berarti kunjungan ke website. Keduanya berhubungan, tetapi mengukur hal yang berbeda: kebutuhan ada pada calon pembeli, sedangkan traffic adalah aktivitas yang terjadi di website. Sebuah kebutuhan juga dapat ada tanpa diungkapkan melalui search, misalnya ketika pembeli langsung bertanya kepada rekan kerja.')
p(mod,'Perbedaan ini penting karena topik yang berhubungan dengan bidang usaha kita belum tentu berasal dari orang yang membutuhkan layanan kita. Artikel tentang cara membersihkan AC sendiri, misalnya, mungkin berguna bagi pemilik rumah. Bila bisnis kita khusus melayani kontrak perawatan AC kantor, kunjungan tersebut belum menunjukkan kebutuhan yang cocok. Sebaliknya, orang yang sedang belajar tentang perawatan AC kantor belum tentu siap membeli sekarang, tetapi pertanyaannya masih dapat menjadi bagian dari keputusan pembelian di kemudian hari.')
p(mod,'SEO, singkatan dari search engine optimization, membantu mesin pencari memahami konten dan membantu orang menemukan website serta menilai apakah mereka perlu mengunjunginya. Google menjelaskan fungsi dasar ini dalam SEO Starter Guide. [S1] Dalam konteks lesson ini, kita menggunakan search visibility untuk menyebut kehadiran atau kemunculan bisnis di hasil pencarian. Kita membatasi pembahasan pada peran pencarian organik, yaitu hasil pencarian yang bukan iklan, dalam perjalanan menuju inquiry.')
p(mod,'Commercial intent berarti tujuan pencarian yang berkaitan dengan mempertimbangkan atau membeli solusi. Pencarian tentang penyedia jasa, cakupan layanan, atau perbandingan penawaran bisa memberi petunjuk ke arah itu. Namun, kata pencarian tidak membuktikan bahwa seseorang mempunyai kebutuhan yang sesuai, wewenang membeli, atau kesiapan untuk melanjutkan. Kita tetap perlu memeriksa konteksnya.')
p(mod,'Jadi, nilai awal sebuah pencarian terletak pada hubungannya dengan kebutuhan yang dapat kita layani. Ini belum menjadi penilaian besarnya pasar atau urutan keyword yang harus dikerjakan. Untuk C01L01, cukup pahami hubungan tersebut sebelum mengejar jumlah kunjungan.')

h(mod,'2 Empat momen sebelum inquiry')
p(mod,'Inquiry adalah pertanyaan atau permintaan yang benar-benar masuk ke bisnis, misalnya pesan tentang layanan atau permintaan penawaran. Sebelum inquiry terjadi, calon pembeli perlu menyelesaikan sejumlah ketidakpastian. Kita memakai empat momen untuk melihatnya: Discover, Learn, Compare/Validate, dan Act. Ini adalah Kultivate framework untuk menyusun pertanyaan, bukan klasifikasi resmi Google.')
table(mod,['Momen','Yang sedang dicari calon pembeli','Peran informasi dari bisnis'],[
 ('Discover','Masalah atau kemungkinan solusi yang sebelumnya belum jelas.','Membantu mengenali masalah dan kaitannya dengan layanan.'),
 ('Learn','Cara kerja solusi dan hal yang perlu dipertimbangkan.','Menjelaskan pilihan dengan bahasa yang dapat dipahami.'),
 ('Compare/Validate','Kecocokan pilihan dan alasan untuk percaya.','Memperjelas cakupan, batas, dan bukti yang bisa diperiksa.'),
 ('Act','Langkah yang diperlukan untuk mulai berinteraksi.','Memperjelas cara menghubungi bisnis dan apa yang terjadi setelahnya.')
],[3.0,6.5,7.3])
p(mod,'Pada Discover, seseorang bisa menyadari bahwa masalahnya perlu ditangani tanpa tahu nama solusi atau penyedianya. Pertanyaan yang kita bantu jawab masih mendasar. Pada Learn, ia mulai memahami pilihan. Penjelasan yang berguna membantu ia membedakan pilihan, termasuk kapan sebuah layanan mungkin belum diperlukan. Karena itu, setiap konten awal tidak harus langsung meminta orang menghubungi sales.')
p(mod,'Pada Compare/Validate, calon pembeli memeriksa apakah pilihan itu cocok dan dapat dipercaya. Compare berarti membandingkan alternatif; Validate berarti memeriksa kembali keyakinan terhadap pilihan. Seseorang yang mendapat nama vendor dari teman dapat menggunakan search untuk memeriksa layanan atau bukti pekerjaannya. Dalam keadaan seperti ini, search membantu validasi meskipun awal perkenalannya berasal dari referral.')
p(mod,'Pada Act, perhatian bergeser ke langkah berikutnya: informasi apa yang perlu disiapkan, siapa yang akan merespons, dan apakah harus memesan pertemuan atau mengirim kebutuhan lebih dahulu. Untuk bisnis jasa, tindakan ini dapat berupa inquiry. Belum ada kepastian bahwa calon pembeli akan menandatangani kontrak.')
p(mod,'Keempat momen ini tidak harus dilalui secara berurutan. Seseorang bisa langsung mencari nama vendor, kembali mempelajari pilihan, lalu meminta pendapat orang lain. Think with Google menjelaskan eksplorasi dan evaluasi yang dapat berulang dalam riset perilaku konsumen. [S4] Temuan itu membantu kita memahami batas gambar alur yang terlalu lurus. Penerapannya ke layanan B2B dalam modul ini tetap merupakan interpretasi Kultivate, bukan kesimpulan bahwa seluruh pembeli B2B mengikuti pola yang sama.')
p(mod,'Gunakan model ini untuk mencari pertanyaan yang belum terjawab. Jangan menganggap empat momen harus menghasilkan empat halaman baru. Satu halaman dapat membantu lebih dari satu momen; penentuan jenis dan susunan halaman dibahas pada lesson berikutnya.')

h(mod,'3 Membaca sinyal sampai menjadi peluang')
p(mod,'Model empat momen menjelaskan apa yang sedang diputuskan calon pembeli. Untuk menilai apakah proses itu menghasilkan manfaat bisnis, kita memerlukan bukti yang berbeda pada setiap bagian. Search visibility, kunjungan yang relevan, inquiry, dan qualified opportunity bukan istilah yang dapat saling menggantikan.')
p(mod,'Qualified opportunity adalah peluang yang sudah diperiksa kecocokannya dan layak dibahas lebih lanjut oleh bisnis. Misalnya, kebutuhan calon pembeli sesuai dengan layanan, lokasinya dapat dilayani, dan ada langkah lanjutan yang masuk akal. Kriteria akhirnya perlu disepakati oleh bisnis sendiri. Inquiry yang hanya menanyakan layanan di luar cakupan belum menjadi peluang yang sesuai. Bahkan peluang yang sudah sesuai pun masih dapat berakhir tanpa penjualan.')
table(mod,['Sinyal','Yang dapat kita pahami','Yang belum terbukti'],[
 ('Visibility','Bisnis atau kontennya muncul di search. Impressions adalah salah satu indikatornya.','Orang tersebut membutuhkan layanan kita.'),
 ('Kunjungan relevan','Konteks kunjungan tampak berkaitan dengan kebutuhan yang dapat dilayani.','Pengunjung sudah berniat atau siap membeli.'),
 ('Inquiry','Ada pesan atau permintaan yang benar-benar diterima.','Kebutuhan dan kondisi calon pembeli sudah cocok.'),
 ('Qualified opportunity','Bisnis telah memeriksa kecocokan dan menyepakati langkah lanjutan.','Penjualan atau revenue sudah terjadi.')
],[3.4,6.3,7.1])
p(mod,'Search Console membedakan impressions, yaitu kemunculan menurut aturan jenis hasil pencarian, dengan clicks, yaitu klik dari Google ke website. [S2] Data ini menjelaskan interaksi di search. Untuk memahami relevansi bisnis, kita perlu membacanya bersama konteks halaman dan pertanyaan yang masuk. Jumlah impressions bukan jumlah calon pembeli unik, dan clicks tidak otomatis sama dengan jumlah inquiry.')
p(mod,'Google Analytics memakai istilah key event untuk tindakan yang dianggap penting bagi bisnis dan direkam sebagai event. [S3] Dalam penerapan kita, klik tombol kontak dapat menjadi sinyal ketertarikan. Tapi klik tersebut belum membuktikan ada pesan yang dikirim dan diterima. Demikian pula, pengiriman formulir yang tercatat perlu dibedakan dari penilaian kualitas kebutuhan oleh tim bisnis. Nama event tidak dapat menggantikan pemeriksaan tersebut.')
p(mod,'Bayangkan sebuah inquiry meminta layanan yang tidak kita tawarkan. Traffic dan inquiry memang terjadi, tetapi kecocokan bisnis belum ada. Sebaliknya, seseorang dapat menemukan penjelasan kita lewat search, kembali melalui tautan yang disimpan, lalu menghubungi lewat kanal lain. Kita tidak dapat mengetahui seluruh perjalanan itu hanya dari kunjungan terakhir. Attribution adalah cara membagikan kontribusi kepada interaksi pemasaran; pembahasan model dan pengukurannya berada di Lesson 06.')
p(mod,'Pada lesson ini, urutan visibility, kunjungan relevan, inquiry, dan qualified opportunity berfungsi sebagai peta pemeriksaan bukti. Diagramnya bukan laporan funnel yang sudah menghubungkan orang yang sama pada setiap tahap. Bila traffic naik tetapi peluang belum jelas, kesimpulan yang tepat adalah memeriksa di mana hubungannya belum terbukti, bukan langsung menyatakan SEO berhasil atau gagal.')

h(mod,'4 Menguji peta pada satu layanan')
p(mod,'Contoh hipotetis. Kita menggunakan penyedia jasa perawatan AC kantor untuk memperlihatkan reasoning, tanpa nama perusahaan, data performa, atau klaim sebagai klien Kultivate. Diasumsikan bisnis ini melayani perawatan berkala untuk kantor di area yang sudah ditetapkan. Seluruh pencarian di bawah adalah ilustrasi, bukan keyword hasil riset atau data pelanggan.')
table(mod,['Momen','Contoh pencarian hipotetis','Jawaban yang dibutuhkan'],[
 ('Discover','AC kantor sering bermasalah','Cara memahami masalah dan kapan perlu bantuan penyedia jasa.'),
 ('Learn','Perawatan AC kantor berkala mencakup apa','Penjelasan lingkup perawatan dan informasi yang perlu disiapkan.'),
 ('Compare/Validate','Vendor perawatan AC kantor dan cakupan layanan','Area layanan, ruang lingkup, batas pekerjaan, serta bukti relevan yang dapat diverifikasi.'),
 ('Act','Permintaan survei perawatan AC kantor','Cara mengirim kebutuhan dan kejelasan proses tindak lanjut.')
],[3.0,6.2,7.6])
p(mod,'Pada momen Learn, calon pembeli masih membutuhkan pemahaman sebelum membandingkan penawaran. Jika halaman hanya berisi ajakan meminta harga, sebagian ketidakpastiannya belum terjawab. Pada Compare/Validate, penjelasan umum juga mungkin belum cukup: ia perlu mengetahui apakah kantor di lokasinya dapat dilayani dan apa yang termasuk dalam pekerjaan. Isi yang diperlukan berubah karena keputusan yang dihadapi berubah.')
p(mod,'Sekarang kita isi satu bagian worksheet. Momen yang dipilih adalah Compare/Validate. Pertanyaan calon pembeli: apakah vendor melayani lokasi kantor dan kebutuhan perawatan kami? Jawaban atau bukti yang perlu tersedia: area layanan dan batas cakupan yang jelas. Dalam skenario ini, saat meninjau halaman, kita menemukan tombol kontak tetapi tidak menemukan informasi area layanan. Itu observasi tentang isi halaman. Dugaan bahwa calon pembeli akhirnya batal menghubungi tetap merupakan hipotesis, yaitu penjelasan sementara yang perlu diperiksa.')
p(mod,'Langkah bukti berikutnya adalah meninjau pertanyaan yang benar-benar diterima tim dan meminta penjelasan dari orang yang menangani inquiry: apakah lokasi layanan memang sering perlu diklarifikasi? Jika ya, informasi tersebut layak diperjelas. Jika tidak, kita perlu melihat kebutuhan lain yang lebih menentukan. Kita belum dapat menyimpulkan bahwa tambahan informasi pasti menaikkan inquiry, apalagi menghitung nilai penjualannya.')
p(mod,'Misalkan kemudian ada permintaan survei yang masuk. Tim masih harus memeriksa lokasi, jenis kebutuhan, dan kesiapan langkah lanjutan. Baru setelah itu peluang dapat dinilai sesuai dengan kriteria bisnis. Logika ini menunjukkan mengapa pencarian yang tampak komersial dan tombol kontak yang terlihat jelas tetap belum menjamin adanya penjualan.')
p(mod,'Untuk menerapkan modelnya, pilih satu layanan utama pada Search-to-Sales Mapping Worksheet bagian C01L01. Jawab empat pertanyaan tentang Discover, Learn, Compare/Validate, dan Act. Gunakan pertanyaan yang pernah didengar dari calon pelanggan bila tersedia. Jika jawabannya masih perkiraan, tandai sebagai hipotesis; jika belum diketahui, tulis belum diketahui. Jangan mengisi kepastian hanya untuk menyelesaikan formulir.')
p(mod,'Setelah itu, pilih satu pertanyaan yang belum terjawab dengan baik. Catat apa yang benar-benar terlihat, apa yang masih kita duga, dan bukti apa yang dapat membedakannya. Peta awal dianggap berguna ketika kita dapat menjelaskan satu gap dan pemeriksaan berikutnya. Ini bukan skor potensi pasar atau rekomendasi membuat konten secara otomatis.')
p(mod,'Search juga belum tentu menjadi prioritas pekerjaan pertama. Jika kita belum tahu apakah calon pembeli memakai search untuk masalah ini, periksa perilaku mereka terlebih dahulu. Jika inquiry yang relevan sudah masuk tetapi respons bisnis tertunda, proses tindak lanjut perlu diperiksa sebelum menambah traffic. Google sendiri tidak menjanjikan ranking pertama atau dampak yang sama untuk setiap perubahan website. [S1]')
p(mod,'Pengetahuan untuk membuat peta awal dan memeriksa pertanyaannya tersedia penuh dalam materi ini. Menentukan besarnya pasar, menyusun dataset keyword, meneliti kompetitor, membuat model attribution, dan menghitung prioritas peluang membutuhkan riset khusus bisnis. Itu merupakan batas pekerjaan layanan pada C01L01. Learner dapat melanjutkannya bersama tim internal atau penyedia jasa sesuai kebutuhan; menggunakan Kultivate bukan syarat untuk memahami atau menerapkan kerangka dasarnya.')
p(mod,'Inti yang dibawa pulang: Nilai search dari keputusan calon pembeli yang dibantunya, lalu periksa bukti menuju inquiry yang sesuai. Traffic saja belum membuktikan peluang bisnis.',lead='Inti yang dibawa pulang:')
p(mod,'Langkah berikutnya adalah Lesson 02, Map Customer Demand Before Choosing Keywords atau Petakan Demand Pelanggan Sebelum Memilih Keyword. Di sana kita mulai menyusun kebutuhan dan pertanyaan pelanggan secara lebih teratur. Worksheet C01L01 dapat dibawa sebagai catatan awal, dan tetap bersifat opsional.')

h(mod,'Referensi dan catatan sumber')
p(mod,'Sumber web berikut diperiksa pada '+DATE+'. Penanda S merujuk fakta atau batas definisi dalam materi. Framework, pertanyaan evaluasi, dan penerapan contoh merupakan reasoning Kultivate. T1 dan T2 dipakai untuk desain pengajaran, bukan untuk membuktikan hasil bisnis.')
for code,org,title,url,note in SOURCES:
    link(mod,f'{code}  {org} — {title}',url); p(mod,note,size=9)
p(mod,'Fondasi internal: Free Course Content and Consultation Funnel Plan (9 September 2026); Content Planning for 6 Courses / 36 Lessons (11 September 2026); Content Production Master Prompt (11 September 2026); Trust Explanation Framework (11 September 2026). Ketentuan format pada brief revisi dan instruksi terbaru menggantikan arahan lama tentang script, cue, serta bentuk ringkasan modul. Topik, urutan, tujuan resource, dan batas layanan tetap mengikuti fondasi tersebut.',size=9)
mod.save(OUT/'C01L01_Modul_Pembelajaran.docx')

ws=docbase('WORKSHEET')
ws.add_heading('Search to Sales Mapping Worksheet',0)
p(ws,'Bagian C01L01    Menghubungkan pencarian dengan peluang bisnis',size=12)
p(ws,'Tujuan kita adalah membuat peta awal untuk satu layanan dan memilih satu bukti yang perlu diperiksa. Gunakan catatan pelanggan atau isi website bila tersedia. Tuliskan hipotesis untuk dugaan dan belum diketahui untuk informasi yang belum ada. Tidak perlu tool berbayar. Sasaran pengerjaan mandiri sekitar 10–15 menit.')
p(ws,'Layanan yang dipilih dan siapa yang dilayani:'); lines(ws,1)
p(ws,'Peran Anda dan URL website bila ada:'); lines(ws,1)
p(ws,'Tujuan bisnis atau masalah organic search saat ini:'); lines(ws,1)
p(ws,'Timeline bila sudah ada: ____________________  Tanggal pencatatan: ____________________',size=9)
h(ws,'Empat pertanyaan calon pembeli')
questions=[
 ('Discover','Apa yang mungkin mereka cari sebelum mengenal layanan atau nama bisnis kita?'),
 ('Learn','Apa yang perlu mereka pahami tentang solusi sebelum menilai kecocokannya?'),
 ('Compare/Validate','Apa yang mereka bandingkan dan bukti apa yang dibutuhkan untuk percaya?'),
 ('Act','Apa yang perlu mereka ketahui agar dapat mengirim kebutuhan atau menghubungi kita?')
]
for stage,q in questions:
    p(ws,stage+' — '+q,lead=stage); lines(ws,2)
p(ws,'Tandai jawaban sebagai observasi, hipotesis, atau belum diketahui. Empat momen ini dapat berulang dan tidak mewajibkan empat halaman baru.',size=9)
ws.add_page_break()
h(ws,'Satu gap dan bukti berikutnya')
p(ws,'Pilih satu momen dari halaman pertama: ____________________________________________')
p(ws,'Pertanyaan calon pembeli yang perlu kita bantu jawab:'); lines(ws,2)
p(ws,'Jawaban atau bukti yang seharusnya tersedia:'); lines(ws,2)
p(ws,'Observasi sekarang dan sumbernya. Tulis apa yang benar-benar terlihat, lokasi halaman atau catatan, serta tanggal pemeriksaan.'); lines(ws,2)
p(ws,'Hipotesis atau hal yang belum diketahui. Pisahkan dari observasi di atas.'); lines(ws,2)
p(ws,'Bukti berikutnya yang perlu diperiksa dan siapa yang dapat membantu:'); lines(ws,2)
p(ws,'Tindakan setelah pemeriksaan dan kapan kita akan meninjau kembali:'); lines(ws,2)
h(ws,'Cara membaca hasil')
p(ws,'Bila pertanyaan masih berasal dari asumsi internal, periksa dulu dengan orang yang menghadapi pelanggan atau catatan inquiry. Bila informasi yang dibutuhkan sudah tersedia, periksa apakah mudah ditemukan dan apakah pertanyaannya memang relevan. Bila gap didukung bukti, catat perbaikan yang masuk akal untuk ditelaah lebih lanjut. Belum ada alasan untuk otomatis membuat halaman baru atau menaikkan anggaran SEO.')
p(ws,'Kelengkapan peta: ____ dari 4 momen sudah mempunyai catatan. Angka ini hanya menunjukkan kelengkapan isian, bukan skor kesehatan SEO atau peluang penjualan.',lead='Kelengkapan peta:',size=9)
p(ws,'Hasil yang kita bawa: satu layanan, empat momen yang dipetakan, satu gap atau unknown, dan satu pemeriksaan berikutnya. Simpan untuk Lesson 02. Estimasi pasar, dataset keyword, riset kompetitor, attribution, dan prioritas peluang membutuhkan pemeriksaan khusus bisnis.',size=9)
ws.save(OUT/'C01L01_Worksheet_Search_to_Sales.docx')

slides=[
 dict(n=1,title='Menghubungkan pencarian\ndengan peluang bisnis',purpose='Menempatkan lesson dalam seri dan mengenalkan satu fokus belajar.',content=['6 Courses × 6 Lessons = 36 Lessons','1 lesson = 1 video utama','Satu lesson, satu keputusan atau mental model','Course 01 • Lesson 01'],visual='Judul editorial, angka struktur seri dalam satu baris, posisi C01L01. Tanpa daftar seluruh topik.',source='Kurikulum internal dan brief revisi. Logo berasal dari file asli pengguna.',chapter='Metadata dan Posisi kita di Kultivate Learn'),
 dict(n=2,title='Traffic naik, peluangnya belum jelas',purpose='Membedakan laporan aktivitas dengan pertanyaan bisnis.',content=['Yang sudah terlihat: website makin banyak dikunjungi','Yang perlu diperiksa: kebutuhan siapa yang terbantu?','Apakah ada inquiry yang sesuai?'],visual='Perbandingan dua bidang terbuka, tanpa grafik atau angka performa.',source='Kultivate framing; kondisi ilustratif, bukan laporan bisnis.',chapter='Bab 1'),
 dict(n=3,title='Demand dan traffic',purpose='Memperjelas kebutuhan calon pembeli dan aktivitas website.',content=['Demand: kebutuhan terhadap solusi','Traffic: kunjungan ke website','Relevansi: apakah kebutuhan itu dapat kita layani?'],visual='Dua definisi besar berdampingan dengan satu pertanyaan penghubung di bawah.',source='Definisi kerja dalam modul dan reasoning Kultivate.',chapter='Bab 1'),
 dict(n=4,title='Empat momen pencarian',purpose='Mengenalkan model utama beserta batas urutannya.',content=['Discover — mengenali masalah','Learn — memahami solusi','Compare/Validate — menilai kecocokan dan bukti','Act — mengambil langkah berikutnya','Momen dapat berulang atau terlewati'],visual='Empat simpul mendatar dengan panah penghubung dan jalur kembali. Label Kultivate framework.',source='Framework kurikulum internal. Batas perjalanan nonlinier merujuk S4, bukan validasi model empat tahap.',chapter='Bab 2'),
 dict(n=5,title='Informasi mengikuti keputusan pembeli',purpose='Menunjukkan pertanyaan yang perlu dibantu pada tiap momen.',content=['Discover: masalah apa yang perlu ditangani?','Learn: bagaimana solusi bekerja?','Compare/Validate: cocok dan bisa dipercaya?','Act: bagaimana memulai?','Satu halaman dapat membantu lebih dari satu momen'],visual='Empat baris berurutan berisi momen dan pertanyaan, tanpa panel bertumpuk.',source='Kultivate framework.',chapter='Bab 2'),
 dict(n=6,title='Bukti menuju peluang bisnis',purpose='Memisahkan sinyal dari proses kualifikasi bisnis.',content=['Visibility','Kunjungan relevan','Inquiry','Qualified opportunity','Setiap perpindahan perlu bukti','Peluang sesuai belum berarti penjualan'],visual='Alur empat tahap. Di bawah setiap tahap terdapat bukti yang diperiksa: kemunculan, konteks kebutuhan, pesan diterima, kecocokan diperiksa.',source='S2 untuk impressions/clicks; kualifikasi dan urutan bukti merupakan Kultivate framework.',chapter='Bab 3'),
 dict(n=7,title='Klik kontak belum membuktikan inquiry',purpose='Menjelaskan mengapa event digital perlu dibedakan dari hasil bisnis.',content=['Klik tombol kontak','Pesan benar-benar diterima','Kebutuhan diperiksa tim','Key event mencatat tindakan penting, bukan otomatis peluang yang sesuai'],visual='Urutan tiga tindakan dengan celah yang menandai pemeriksaan. Label contoh ilustratif.',source='S3 untuk definisi key event. Penerapan tombol kontak merupakan ilustrasi Kultivate.',chapter='Bab 3'),
 dict(n=8,title='Contoh perawatan AC kantor',purpose='Membuat Discover dan Learn konkret dalam satu skenario B2B.',content=['Discover: AC kantor sering bermasalah','Perlu memahami masalah','Learn: perawatan berkala mencakup apa','Perlu memahami lingkup layanan','Contoh hipotetis, bukan keyword hasil riset'],visual='Dua langkah besar bernomor dengan pencarian di atas dan kebutuhan informasi di bawah. Tidak menyerupai hasil Google nyata.',source='Contoh hipotetis Kultivate. Tidak memakai data atau nama klien.',chapter='Bab 4'),
 dict(n=9,title='Kecocokan diperiksa sebelum menjadi peluang',purpose='Menyelesaikan contoh dengan Compare/Validate dan Act.',content=['Compare/Validate: area layanan, cakupan, bukti relevan','Act: kirim kebutuhan, pahami tindak lanjut','Permintaan masuk masih perlu diperiksa kecocokannya','Contoh hipotetis'],visual='Dua bagian berurutan lalu satu kesimpulan pendek. Struktur melanjutkan contoh slide 8.',source='Contoh hipotetis Kultivate.',chapter='Bab 4'),
 dict(n=10,title='Observasi dan hipotesis',purpose='Menunjukkan cara mengisi gap worksheet tanpa mengarang diagnosis.',content=['Observasi: area layanan tidak tercantum di halaman','Hipotesis: calon pembeli ragu lalu batal menghubungi','Bukti berikutnya: periksa pertanyaan lokasi dalam inquiry','Contoh hipotetis'],visual='Tiga baris besar dengan label yang jelas dan garis pemisah. Hubungan observasi, dugaan, pemeriksaan terlihat.',source='Penerapan hypothetical worked example Kultivate.',chapter='Bab 4'),
 dict(n=11,title='Peta awal untuk layanan Anda',purpose='Menghubungkan lesson dengan worksheet dan batas tindakan.',content=['Pilih satu layanan','Jawab empat pertanyaan pembeli','Catat satu gap atau unknown','Tentukan bukti berikutnya','Belum tahu perilaku search? Periksa dahulu.','Inquiry relevan terlambat ditangani? Periksa tindak lanjut.'],visual='Alur isian di sisi kiri, dua kondisi keputusan di sisi kanan. Resource disebut dengan nama yang sama.',source='Search-to-Sales Mapping Worksheet bagian C01L01; Kultivate decision framework.',chapter='Bab 4 dan Worksheet'),
 dict(n=12,title='Search dinilai lewat keputusan yang dibantunya',purpose='Menutup dengan satu takeaway dan jalur belajar berikutnya.',content=['Periksa bukti menuju inquiry yang sesuai.','Traffic saja belum membuktikan peluang bisnis.','Berikutnya: Lesson 02','Map Customer Demand Before Choosing Keywords','Worksheet C01L01 bersifat opsional'],visual='Takeaway besar dengan next lesson yang lebih kecil. Tanpa CTA konsultasi utama.',source='Sintesis Kultivate, sesuai CTA kurikulum untuk lesson awal.',chapter='Bab 4 bagian penutup')
]
(BUILD/'slides.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
(BUILD/'sources.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2),encoding='utf-8')

plan=docbase('SLIDE PLANNING DAN QUALITY CHECK')
plan.add_heading('Rencana slide dan pemeriksaan sample',0)
p(plan,'Course 01 Lesson 01    Menghubungkan pencarian dengan peluang bisnis',size=12)
p(plan,'Paket sample berisi metadata dan modul pembelajaran, bagian worksheet C01L01, deck presentasi aktual, serta dokumen perencanaan dan pemeriksaan ini. Statusnya menunggu persetujuan format. Lesson 02–36 belum dikerjakan.')
h(plan,'Keputusan editorial dan pengajaran')
p(plan,'Modul dibangun sebagai bahan pemahaman pengajar dengan empat chapter utama. Alurnya bergerak dari kebutuhan calon pembeli, empat momen pencarian, bukti menuju peluang, lalu satu contoh dan aplikasi. Paragraf menjadi bentuk utama. Tabel hanya dipakai ketika perbandingan antar momen atau sinyal lebih mudah dibaca bersama.')
p(plan,'Halaman Improvise Acquisition.com ditinjau ulang untuk progression bertahap serta kedekatan resources dan personalized help. Halaman SEO Unlocked ditinjau ulang untuk phase, week, lesson, itinerary, dan worksheet. Adaptasinya pada sample adalah posisi belajar yang jelas, model yang ringkas, contoh yang selesai dijelaskan, dan resource untuk aplikasi. Peninjauan ini dilakukan atas halaman publik, bukan transkrip seluruh video. Tidak ada klaim bahwa format tersebut terbukti menaikkan conversion Kultivate.')
p(plan,'Dua belas slide dipilih berdasarkan fungsi penjelasannya: orientasi, masalah, definisi, model, pertanyaan pembeli, bukti, satu pembedaan event, dua bagian contoh, interpretasi worksheet, aplikasi, dan takeaway. Deck ditujukan untuk satu video 6–12 menit; durasi nyata tetap bergantung pada penjelasan Hans. Tidak ada timestamp, instruksi presenter, atau naskah kata per kata.')
p(plan,'Arah visual mengikuti Dark Editorial Kultivate dengan Void Black, Cloud White, Mist, dan aksen Cyan. Diagram dan teks berupa elemen native yang dapat diedit. Arial digunakan karena tersedia dalam lingkungan rendering dan PowerPoint saat ini; pilihan font merupakan adaptasi produksi, bukan perubahan spesifikasi font brand. Logo memakai file asli pengguna tanpa digambar ulang. Perubahan komposisi mengikuti materi, bukan satu layout berulang.')
h(plan,'Peta fondasi yang dipertahankan')
table(plan,['Fondasi','Implementasi pada sample'],[
 ('Course 01 dan Lesson 01','Urutan dan judul Inggris tetap. Judul Indonesia mengikuti brief terbaru.'),
 ('Tujuan pembelajaran','Discovery, evaluasi, keputusan sebelum inquiry; visibility dihubungkan ke kebutuhan yang sesuai.'),
 ('Teaching spine','Discover, Learn, Compare/Validate, Act tetap menjadi model utama.'),
 ('Search-to-Sales Mapping Worksheet','Tetap satu resource tingkat course. Sample hanya mengisi bagian C01L01.'),
 ('Batas layanan','Market sizing, dataset keyword, riset kompetitor, attribution model, prioritas estimasi peluang tetap pekerjaan khusus bisnis.'),
 ('CTA awal course','Lesson 02 menjadi next step utama; worksheet opsional.'),
 ('Instruksi script lama','Digantikan oleh modul pembelajaran dan deck. Tidak ada deliverable script dalam paket baru.')
],[5,11.8])
p(plan,'Catatan ketersediaan sumber: file yang ditemukan berisi rencana kurikulum, master prompt, dan research, bukan modul C01L01 lengkap atau worksheet terisi yang sudah final. Bagian worksheet dikembangkan dari nama, tujuan, model, dan ketentuan resource yang tercantum di fondasi. Tambahan field observasi, hipotesis, dan bukti berikutnya memperjelas aplikasi, tanpa mengganti tujuan resource. Ukuran kelengkapan isian menggantikan kebutuhan skor performa yang tidak memiliki dasar pada lesson ini.',size=9)
plan.add_page_break()
h(plan,'Rencana setiap slide')
for s in slides:
    plan.add_heading(f"Slide {s['n']:02d} {s['title'].replace(chr(10),' ')}",2)
    p(plan,'Purpose: '+s['purpose'],lead='Purpose:',size=10)
    p(plan,'On-slide content: '+' / '.join(s['content']),lead='On-slide content:',size=10)
    p(plan,'Suggested visual: '+s['visual'],lead='Suggested visual:',size=10)
    p(plan,'Source: '+s['source'],lead='Source:',size=9)
    p(plan,'Modul yang didukung: '+s['chapter'],lead='Modul yang didukung:',size=9)

plan.add_page_break()
h(plan,'Research dan claim log')
p(plan,'Tanggal pemeriksaan web: '+DATE+'. Klaim fakta dibuat terbatas pada isi sumber. Interpretasi dan rekomendasi tidak diperlakukan sebagai fakta industri.')
for code,org,title,url,note in SOURCES:
    link(plan,f'{code}  {org} — {title}',url); p(plan,note,size=9)
table(plan,['Jenis','Pernyataan dan batas pemakaian'],[
 ('Fakta sumber','S1: SEO membantu pemahaman dan penemuan konten; tidak menjamin ranking pertama. Tidak digunakan untuk menjanjikan inquiry.'),
 ('Fakta sumber','S2: impressions dan clicks berbeda; aturan impressions mengikuti jenis hasil. Tidak dipakai sebagai jumlah pembeli unik.'),
 ('Fakta sumber','S3: key event mencatat tindakan penting. Kualitas peluang bukan penilaian otomatis dari label event.'),
 ('Fakta dan batas','S4: eksplorasi dan evaluasi dapat berulang dalam riset konsumen. Tidak digeneralisasi sebagai perilaku wajib semua pembeli B2B Indonesia.'),
 ('Kultivate framework','Empat momen dan peta bukti adalah alat reasoning. Bukan model resmi Google, causal model, atau funnel terukur.'),
 ('Ilustrasi','Seluruh pencarian, halaman, inquiry, dan dugaan pada contoh AC kantor bersifat hipotetis. Tidak ada angka performa, perusahaan palsu, atau testimonial.'),
 ('Interpretasi','Gap informasi mungkin menghambat keputusan, tetapi perlu bukti pelanggan sebelum dipakai sebagai diagnosis.'),
 ('Keputusan kurikulum','6 × 6 adalah pilihan desain Kultivate. Bukan angka optimal yang diklaim dari riset.')
],[3.5,13.3])
h(plan,'Quality check isi dan bahasa')
p(plan,'Content: Definisi demand, traffic, SEO, commercial intent, inquiry, qualified opportunity, key event, dan attribution dijelaskan saat diperlukan. Model, contoh, aplikasi, dan batas kesimpulan tersedia. Detail keyword, audit halaman, implementasi analytics, dan estimasi pasar tetap di luar lesson ini.')
p(plan,'Language: Bahasa Indonesia menjadi bahasa utama dengan sudut pandang kita. Istilah Inggris disertai arti dan konteks. Tidak ada dialog, opening atau closing script, cue produksi, hype, atau klaim pengalaman klien.')
p(plan,'Structure: Empat chapter utama mempunyai perubahan reasoning yang jelas. Tabel membandingkan momen, sinyal, dan contoh. Hanya ada satu core takeaway. Orientasi seri, metadata, dan referensi berada di luar chapter sebagai kebutuhan navigasi.')
p(plan,'Trust: Sumber fakta dapat dilacak melalui tautan. Diagram diberi label Kultivate framework dan contoh diberi label hipotetis. Fakta, dugaan, serta tindakan pemeriksaan dibedakan. CTA utama adalah melanjutkan belajar, tanpa mengurangi pengetahuan inti.')
p(plan,'Worksheet: Dua halaman, satu layanan, empat pertanyaan keputusan, satu gap, observasi, hipotesis atau unknown, bukti berikutnya, dan cara membaca hasil. Tidak memerlukan proprietary tool atau pengiriman data untuk digunakan. Tidak ada skor SEO atau peluang penjualan yang dibuat-buat.')
h(plan,'Pemeriksaan visual dan file')
p(plan,'Hasil pemeriksaan file dan render final dicatat di bawah setelah proses ekspor. Pemeriksaan ini mencakup keterbacaan setiap halaman, kecocokan planning dengan slide aktual, teks terpotong atau bertumpuk, label sumber dan contoh, serta keberadaan elemen editable. Catatan ini bukan bagian dari slide pengajaran.')
p(plan,'QA_RENDER_STATUS')
h(plan,'Batas persetujuan sample')
p(plan,'Persetujuan sample mencakup kedalaman penjelasan, bahasa, hubungan worksheet, dan arah visual slide. Sample menjadi quality benchmark setelah disetujui; struktur chapter serta bentuk visual tetap dapat mengikuti kebutuhan tiap lesson. Penerapan ke Lesson 02–36 menunggu persetujuan pengguna.')
plan.save(OUT/'C01L01_Slide_Planning_dan_Quality_Check.docx')
print(json.dumps({'outputs':[str(x) for x in OUT.glob('*.docx')],'slides':len(slides)},ensure_ascii=False))
